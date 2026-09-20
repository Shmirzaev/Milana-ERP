from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.models import AttendanceEvent, Employee
from app.tests.conftest import TestSessionLocal
from app.tests.test_attendance import INTEGRATION_HEADERS, person, snapshot


DAY = "2026-08-17"


def _event(uid: str, employee_no: str, time: str, result=...):
    event = {
        "event_uid": uid,
        "external_person_id": employee_no,
        "occurred_at": f"{DAY}T{time}+05:00",
        "direction": "unknown",
        "verification_mode": "face",
    }
    if result is not ...:
        event["result"] = result
    return event


def _import_people_and_events(client, people, events) -> None:
    roster = client.post(
        "/api/attendance/integration/people",
        headers=INTEGRATION_HEADERS,
        json=snapshot(*people),
    )
    assert roster.status_code == 200, roster.text
    imported = client.post(
        "/api/attendance/integration/events",
        headers=INTEGRATION_HEADERS,
        json={"device": snapshot()["device"], "events": events},
    )
    assert imported.status_code == 200, imported.text


def _as_utc(value: str) -> datetime:
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    return parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed.astimezone(timezone.utc)


def test_explicit_failed_and_unknown_events_stay_raw_but_do_not_count_as_attendance(
    client,
    auth_headers,
):
    _import_people_and_events(
        client,
        [person("990001", "Mixed result"), person("990002", "Failed only")],
        [
            _event("mixed-failed", "990001", "07:00:00", "failed"),
            _event("mixed-success", "990001", "08:00:00", " SUCCESS "),
            _event("mixed-unknown", "990001", "19:00:00", "unknown"),
            _event("failed-only", "990002", "09:00:00", "failed"),
        ],
    )

    overview = client.get(f"/api/attendance/overview?day={DAY}", headers=auth_headers)

    assert overview.status_code == 200, overview.text
    body = overview.json()
    rows = {row["external_person_id"]: row for row in body["people"]}
    assert body["summary"]["events_today"] == 4
    assert body["summary"]["total_people"] == 2
    assert body["summary"]["used_today"] == 1
    assert body["summary"]["not_used_today"] == 1
    assert rows["990001"]["event_count"] == 1
    assert _as_utc(rows["990001"]["arrival_at"]) == datetime(
        2026, 8, 17, 3, tzinfo=timezone.utc,
    )
    assert rows["990001"]["attendance_status"] == "single_scan"
    assert rows["990002"]["event_count"] == 0
    assert rows["990002"]["attendance_status"] == "absent"
    used = client.get(f"/api/attendance/overview?day={DAY}&usage=used", headers=auth_headers)
    not_used = client.get(f"/api/attendance/overview?day={DAY}&usage=not_used", headers=auth_headers)
    assert [row["external_person_id"] for row in used.json()["people"]] == ["990001"]
    assert [row["external_person_id"] for row in not_used.json()["people"]] == ["990002"]
    report = client.get(
        f"/api/attendance/reports/daily.xlsx?day={DAY}&lang=en",
        headers=auth_headers,
    )
    assert report.status_code == 200, report.text
    sheet = load_workbook(BytesIO(report.content)).active
    report_rows = {
        sheet.cell(row, 2).value: [sheet.cell(row, column).value for column in range(4, 8)]
        for row in range(7, sheet.max_row + 1)
    }
    assert report_rows["990001"] == ["One scan only", "08:00", None, None]
    assert report_rows["990002"] == ["No scans", None, None, None]
    with TestSessionLocal() as db:
        assert db.query(AttendanceEvent).count() == 4


def test_missing_result_remains_accepted_for_legacy_integration(client, auth_headers):
    _import_people_and_events(
        client,
        [person("990003", "Legacy result")],
        [
            _event("legacy-arrival", "990003", "08:15:00"),
            _event("legacy-departure", "990003", "18:20:00"),
        ],
    )

    overview = client.get(f"/api/attendance/overview?day={DAY}", headers=auth_headers)

    assert overview.status_code == 200, overview.text
    row = overview.json()["people"][0]
    assert row["event_count"] == 2
    assert row["worked_minutes"] == 605
    assert row["attendance_status"] == "complete"


def test_hr_attendance_uses_only_accepted_event_results(client, auth_headers):
    with TestSessionLocal.begin() as db:
        db.add(Employee(
            factory_code="MIL",
            employee_no="990004",
            full_name="HR result policy",
            status="active",
            hr_profile_json={},
        ))
    _import_people_and_events(
        client,
        [person("990004", "HR result policy")],
        [
            _event("hr-failed", "990004", "07:00:00", "failed"),
            _event("hr-success", "990004", "08:00:00", "success"),
            _event("hr-unknown", "990004", "19:00:00", "unknown"),
        ],
    )

    response = client.get(f"/api/hr/attendance?day={DAY}", headers=auth_headers)

    assert response.status_code == 200, response.text
    row = next(item for item in response.json()["rows"] if item["employee_no"] == "990004")
    assert _as_utc(row["arrival_at"]) == datetime(2026, 8, 17, 3, tzinfo=timezone.utc)
    assert row["departure_at"] is None
    assert row["worked_minutes"] == 0
