from io import BytesIO

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.models import LegacyStockReceipt, Model, ModelColor, Package
from app.services.warehouse_report import export_warehouse_report


def test_report_preserves_model_grain_exact_names_colors_and_ready_totals():
    with SessionLocal() as db:
        model = Model(code="REPORT-007", name='  =Название "Точное"\nВторая строка  ',
                      details_json={"general": {"model_no": "REPORT", "variant_no": "007"}})
        unknown = Model(code="REPORT-008", name="Длинное название изделия " * 10)
        db.add_all([model, unknown])
        db.flush()
        db.add_all([ModelColor(model_id=model.id, color_name=value) for value in ['Синий', 'Красный ', 'Синий']])
        for index, (item, kind, status, quantity) in enumerate([
            (model, 'standard', 'packed', 12), (model, 'standard', 'reserved', 8),
            (unknown, 'standard', 'received_in_storage', 5), (model, 'first_grade', 'received_in_storage', 1),
            (model, 'standard', 'shipped', 99), (model, 'standard', 'damaged', 77),
        ]):
            receipt = LegacyStockReceipt(source_system="REPORT_TEST", source_warehouse_id="1", source_record_id=str(index),
                                         source_checksum="0" * 64, source_payload={})
            db.add(receipt)
            db.flush()
            db.add(Package(package_no=f'REPORT-{index}', barcode=f'REPORT-QR-{index}', model_id=item.id,
                           legacy_receipt_id=receipt.id,
                           color='Package color must not replace variant catalog colors', stock_kind=kind,
                           status=status, capacity=quantity, total_quantity=quantity))
        db.flush()
        workbook = load_workbook(BytesIO(export_warehouse_report(db, 'ru')))
        standard = workbook['Основной склад']
        rows = [row for row in standard.iter_rows(min_row=6, max_row=standard.max_row-1, values_only=True) if row[0] == 'REPORT']
        assert rows == [('REPORT', model.name, '007', 'Синий\nКрасный ', 2, 20),
                        ('REPORT', unknown.name, '008', '—', 1, 5)]
        assert standard['B6'].data_type == 's'
        for row in standard.iter_rows(min_row=6, max_row=standard.max_row-1):
            if row[0].value == 'REPORT':
                assert row[2].number_format == '@'
                assert row[1].alignment.wrap_text and row[4].alignment.horizontal == 'right'
                assert row[1].font.name == 'Calibri' and row[1].border.bottom.style == 'thin'
                if row[2].value == '008':
                    assert standard.row_dimensions[row[1].row].height > 26
        assert standard.freeze_panes == 'A6'
        assert standard.auto_filter.ref == f'A5:F{standard.max_row-1}'
        assert standard.cell(standard.max_row, 5).value == f'=SUM(E6:E{standard.max_row-1})'
        assert standard.cell(standard.max_row, 6).value == f'=SUM(F6:F{standard.max_row-1})'
        assert list(standard.values)[4] == ('Номер модели','Название товара','Номер варианта','Цвет','Упаковки','Штуки')
        single = workbook['Первый сорт']
        row = next(row for row in single.iter_rows(min_row=6, max_row=single.max_row-1, values_only=True) if row[0] == 'REPORT')
        assert row == ('REPORT', model.name, '007', 'Синий\nКрасный ', 1, 1)


def test_empty_report_retains_source_headers_and_zero_totals():
    with SessionLocal() as db:
        # The baseline fixture contains no ready packages.
        db.query(Package).delete(synchronize_session=False)
        for lang, title in [('ru', 'Основной склад'), ('en', 'Standard stock'), ('uz', 'Asosiy ombor')]:
            sheet = load_workbook(BytesIO(export_warehouse_report(db, lang)))[title]
            assert sheet.max_row == 6 and sheet.max_column == 6
            assert sheet['E6'].value == sheet['F6'].value == 0
            assert 'UTC+05:00' in sheet['A2'].value and title in sheet['A3'].value
            assert sheet.auto_filter.ref == 'A5:F5'
