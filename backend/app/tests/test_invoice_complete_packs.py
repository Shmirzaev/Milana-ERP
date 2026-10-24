"""Complete-pack aggregation is presentation only, including old snapshots."""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.api.routes.shipments import _invoice_print_details
from app.db.session import SessionLocal
from app.models import LegacyStockReceipt, Package, PackageItem, Shipment, ShipmentPackage
from app.models.shipment_invoice_service import build_invoice_rows, render_shipment_invoice
from app.models.shipment_invoice_excel_service import shipment_invoice_workbook
from app.models.shipment_review_service import shipment_document
from app.tests.test_shipment_review import dispatch as dispatch
from app.tests.test_manual_pack_dispatch import receive, shipment as manual_shipment
from app.tests.test_package_workflows import warehouse as warehouse, package_qr


def invoice_fixture():
    packages = [
        {"package_no": "A", "quantity": 90, "weight_kg": "23.14", "is_partial": False},
        {"package_no": "PART", "quantity": 15, "weight_kg": "4.00", "is_partial": True},
        {"package_no": "B", "quantity": 90, "weight_kg": "23.10", "is_partial": False},
    ]
    lines = [
        {"package_no": pack["package_no"], "model_no": "BM8010", "variant_no": "V-6451",
         "description": "Trousers", "color": "Blue", "size": size, "quantity": 15,
         "unit_price": "2.00", "amount": "30.00"}
        for pack in packages for size in (["48"] if pack["is_partial"] else ["48", "50", "52", "54", "56", "58"])
    ]
    return {"shipment_no": "EXAMPLE", "package_details": packages, "lines": lines,
            "packages_count": 3, "quantity": 195, "amount": "390.00", "invoice_layout_version": 2}


@pytest.mark.parametrize("lang", ["en", "ru", "uz"])
@pytest.mark.parametrize("prices", [True, False])
def test_complete_packs_combine_across_partial_with_exact_print_excel_totals(lang, prices):
    document = invoice_fixture()
    before = deepcopy(document)
    rows = build_invoice_rows(document["lines"], document["package_details"])
    assert len(rows) == 2
    full, partial = rows
    assert (full["pack_count"], full["quantity"], full["weight_kg"], full["amount"]) == (2, 180, "46.24", "360.00")
    assert [s["quantity"] for s in full["sizes"]] == [30] * 6
    assert partial["package_no"] == "PART" and partial["quantity"] == 15 and partial["pack_count"] == 1
    html = render_shipment_invoice(document, lang, show_prices=prices)
    assert html.count('<td>V-6451</td>') == 2
    assert '>23.12</td>' in html and '>46.24</td>' in html and '48×30' in html
    sheet = load_workbook(BytesIO(shipment_invoice_workbook(document, lang, show_prices=prices))).active
    assert [sheet.cell(10, col).value for col in (6, 7, 8, 9)] == [2, 180, 23.12, 46.24]
    assert [sheet.cell(11, col).value for col in (6, 7)] == [1, 15]
    assert [sheet.cell(12, col).value for col in (6, 7, 9)] == [3, 195, 50.24]
    assert sheet["K10"].value == (360 if prices else None)
    assert document == before


@pytest.mark.parametrize("field,value", [("variant_no", "V-9999"), ("model_no", "OTHER"), ("unit_price", "3.00")])
def test_distinct_variant_model_or_price_never_combines(field, value):
    document = invoice_fixture()
    for line in document["lines"]:
        if line["package_no"] == "B":
            line[field] = value
    assert len(build_invoice_rows(document["lines"], document["package_details"])) == 3


def test_missing_weight_price_and_completeness_are_not_invented():
    document = invoice_fixture()
    document["package_details"][2]["weight_kg"] = None
    for line in document["lines"]:
        line["unit_price"] = line["amount"] = None
    rows = build_invoice_rows(document["lines"], document["package_details"])
    assert rows[0]["pack_count"] == 2 and rows[0]["weight_kg"] is None and rows[0]["amount"] is None
    del document["package_details"][2]["is_partial"]
    assert len(build_invoice_rows(document["lines"], document["package_details"])) == 3


def test_aggregate_size_labels_are_not_repeated():
    document = invoice_fixture()
    document["lines"] = [{**line, "size_display": "48–58", "quantity": 90, "amount": "180.00"}
                         for line in document["lines"] if line["size"] == "48" and line["package_no"] != "PART"]
    document["package_details"] = [p for p in document["package_details"] if p["package_no"] != "PART"]
    row = build_invoice_rows(document["lines"], document["package_details"])[0]
    assert row["quantity"] == 180 and len(row["sizes"]) == 1
    assert row["sizes"][0]["aggregate"] is True


def test_draft_and_historical_snapshot_group_without_mutating_frozen_document(client, auth_headers, dispatch):
    with SessionLocal() as db:
        first = db.get(Package, dispatch["package"])
        first.capacity = 8
        first.weight_kg = Decimal("2.35")
        receipt = LegacyStockReceipt(source_system="TEST", source_warehouse_id="invoice", source_record_id="second",
                                     source_checksum="b" * 64, source_payload={"quantity": 8})
        db.add(receipt)
        db.flush()
        second = Package(package_no="PACK-TWO", barcode="PACK-TWO", model_id=first.model_id, color=first.color,
                         capacity=8, total_quantity=8, weight_kg=Decimal("2.45"), legacy_receipt_id=receipt.id, status="reserved")
        db.add(second)
        db.flush()
        db.add(ShipmentPackage(shipment_id=dispatch["shipment"], package_id=second.id, quantity=8))
        for size in ("48", "50"):
            db.add(PackageItem(package_id=second.id, model_id=first.model_id, color=first.color, size=size, quantity=4))
        db.flush()
        shipment = db.get(Shipment, dispatch["shipment"])
        document = shipment_document(db, shipment)
        assert len(document["invoice_rows"]) == 1 and document["invoice_rows"][0]["pack_count"] == 2
        # Legacy records predate is_partial; their saved values remain untouched.
        for pack in document["package_details"]:
            pack.pop("is_partial")
        frozen = deepcopy(document)
        shipment.status = "shipped"
        shipment.dispatch_snapshot = {"document": frozen}
        db.commit()
        presented = _invoice_print_details(db, shipment, dict(frozen))
        assert len(presented["invoice_rows"]) == 1
        assert Decimal(presented["invoice_rows"][0]["weight_kg"]) == Decimal("4.80")
        assert shipment.dispatch_snapshot["document"] == frozen
    base = f'/api/shipments/{dispatch["shipment"]}'
    for suffix in ("/invoice", "/invoice/print", "/invoice.xlsx"):
        assert client.get(base + suffix, headers=auth_headers).status_code == 200
        assert client.get(base + suffix).status_code == 401
    with SessionLocal() as db:
        assert db.get(Shipment, dispatch["shipment"]).dispatch_snapshot["document"] == frozen


@pytest.mark.parametrize("capacity,selected,partial", [(8, 8, False), (10, 8, True), (8, 4, True)])
def test_completeness_uses_pack_capacity_and_selected_quantity(client, dispatch, capacity, selected, partial):
    with SessionLocal() as db:
        package = db.get(Package, dispatch["package"])
        package.capacity = capacity
        shipment = db.get(Shipment, dispatch["shipment"])
        link = db.query(ShipmentPackage).filter_by(shipment_id=shipment.id, package_id=package.id).one()
        link.quantity = selected
        if selected < 8:
            shipment.dispatch_snapshot = {"package_quantities": {str(package.id): {str(i): 2 for i in dispatch["items"]}}}
        db.flush()
        document = shipment_document(db, shipment)
        assert document["package_details"][0]["is_partial"] is partial


def test_receipt_scan_partial_selection_dispatch_and_reprint(client, warehouse):
    run = receive(client, warehouse, (90, 90, 90))
    sid = manual_shipment(client, warehouse, Decimal("2.00"))
    base = f"/api/shipments/{sid}"
    for pid in run["package_ids"]:
        response = client.post(base + "/scan-package", headers=warehouse, json={"code": package_qr(pid)})
        assert response.status_code == 200 and response.json()["ok"]
    pid = run["package_ids"][-1]
    with SessionLocal() as db:
        item = db.query(PackageItem).filter_by(package_id=pid).one()
        item_id = item.id
    response = client.post(base + f"/packages/{pid}/quantity", headers=warehouse,
                           json={"expected_quantity": 90, "keep_remainder": True, "reason": "Customer selected 15 pieces",
                                 "items": [{"item_id": item_id, "quantity": 15}]})
    assert response.status_code == 200, response.text
    for dispatched in (False, True):
        if dispatched:
            response = client.post(base + "/ship", headers=warehouse)
            assert response.status_code == 200, response.text
        document = client.get(base + "/invoice", headers=warehouse).json()
        assert [(r["pack_count"], r["quantity"]) for r in document["invoice_rows"]] == [(2, 180), (1, 15)]
        assert Decimal(document["amount"]) == Decimal("390.00")
        sheet = load_workbook(BytesIO(client.get(base + "/invoice.xlsx", headers=warehouse).content)).active
        assert sheet["F10"].value == 2 and sheet["G10"].value == 180 and sheet["G11"].value == 15
    from app.models import FinishedGoodsStock
    with SessionLocal() as db:
        assert db.get(Package, pid).total_quantity == 90
        assert db.query(FinishedGoodsStock).filter_by(package_id=pid).one().available_qty == 75
