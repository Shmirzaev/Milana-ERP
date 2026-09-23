from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.services import attendance_reports


def test_daily_attendance_export_reuses_alignment_styles_and_preserves_cells(monkeypatch):
    rows = [
        {
            "external_person_id": str(index),
            "full_name": f"Employee {index}",
            "attendance_status": status,
            "arrival_at": None,
            "departure_at": None,
            "worked_minutes": None,
        }
        for index, status in enumerate(
            ["complete", "single_scan", "absent"] * 20,
            start=1,
        )
    ]
    real_alignment = attendance_reports.Alignment
    alignment_calls = []

    def track_alignment(*args, **kwargs):
        alignment_calls.append((args, kwargs))
        return real_alignment(*args, **kwargs)

    monkeypatch.setattr(attendance_reports, "Alignment", track_alignment)
    content = attendance_reports.build_daily_attendance_xlsx(
        day=date(2026, 9, 23),
        rows=rows,
        generated_at=datetime(2026, 9, 23, 3, 0, tzinfo=timezone.utc),
        lang="en",
    )

    sheet = load_workbook(BytesIO(content), data_only=False).active
    assert len(alignment_calls) == 3
    assert sheet["A1"].value == "Daily attendance report"
    assert sheet["A3"].value == "Profiles: 60  Complete: 20  One scan: 20  Absent: 20"
    assert [sheet.cell(7, column).value for column in range(1, 8)] == [
        1,
        "1",
        "Employee 1",
        "Arrival and departure",
        None,
        None,
        None,
    ]
    assert sheet["A6"].alignment.wrap_text is True
    assert sheet["C7"].alignment.horizontal == "left"
    assert sheet["A7"].alignment.horizontal == "center"
