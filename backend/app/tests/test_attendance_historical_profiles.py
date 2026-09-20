from io import BytesIO

from openpyxl import load_workbook

from app.models import AttendancePerson
from app.tests.conftest import TestSessionLocal
from app.tests.test_attendance import INTEGRATION_HEADERS, person, snapshot


def test_removed_profile_remains_in_historical_overview_and_export(client, auth_headers):
    roster = snapshot(person("735", "Historical worker"), person("371", "No scans"))
    assert client.post("/api/attendance/integration/people", headers=INTEGRATION_HEADERS, json=roster).status_code == 200
    events = [{"event_uid": f"historical-{index}", "external_person_id": "735",
               "occurred_at": timestamp, "result": "success"}
              for index, timestamp in enumerate(["2026-08-17T08:00:00+05:00", "2026-08-17T17:00:00+05:00"])]
    assert client.post("/api/attendance/integration/events", headers=INTEGRATION_HEADERS,
                       json={"device": roster["device"], "events": events}).status_code == 200
    assert client.post("/api/attendance/integration/people", headers=INTEGRATION_HEADERS, json=snapshot()).status_code == 200

    response = client.get("/api/attendance/overview?day=2026-08-17", headers=auth_headers)
    assert response.status_code == 200, response.text
    data = response.json()
    assert [row["external_person_id"] for row in data["people"]] == ["735"]
    assert data["people"][0]["worked_minutes"] == 540
    assert data["summary"]["total_people"] == data["summary"]["used_today"] == 1
    assert data["pagination"]["total"] == 1
    report = client.get("/api/attendance/reports/daily.xlsx?day=2026-08-17&lang=en", headers=auth_headers)
    assert report.status_code == 200
    sheet = load_workbook(BytesIO(report.content)).active
    assert sheet.cell(7, 2).value == "735"
    for suffix in ("day=2026-08-18", "day=2026-08-17&usage=not_used", "day=2026-08-17&query=missing"):
        filtered = client.get(f"/api/attendance/overview?{suffix}", headers=auth_headers)
        assert filtered.status_code == 200
        assert filtered.json()["people"] == []


def test_historical_rows_prefer_active_profile_and_keep_factory_scope(client, auth_headers):
    from datetime import datetime, timezone
    from app.models import AttendanceDevice, AttendanceEvent

    timestamp = datetime(2026, 8, 17, 3, tzinfo=timezone.utc)
    with TestSessionLocal.begin() as db:
        first = AttendanceDevice(factory_code="MIL", device_key="history-first", name="First", read_only=True)
        second = AttendanceDevice(factory_code="MIL", device_key="history-second", name="Second", read_only=True)
        other = AttendanceDevice(factory_code="ECO", device_key="history-other", name="Other", read_only=True)
        db.add_all([first, second, other])
        db.flush()
        db.add_all([
            AttendancePerson(factory_code="MIL", device_id=first.id, external_person_id="735", full_name="Old name", present_on_device=False, last_synced_at=timestamp),
            AttendancePerson(factory_code="MIL", device_id=second.id, external_person_id="735", full_name="Current name", present_on_device=True, last_synced_at=timestamp),
            AttendancePerson(factory_code="ECO", device_id=other.id, external_person_id="735", full_name="Other factory", present_on_device=True, last_synced_at=timestamp),
            AttendanceEvent(factory_code="MIL", device_id=first.id, event_uid="historical-scope", external_person_id="735", occurred_at=timestamp, received_at=timestamp, result="success"),
        ])
    result = client.get("/api/attendance/overview?day=2026-08-17", headers=auth_headers)
    assert result.status_code == 200, result.text
    rows = result.json()["people"]
    assert len(rows) == 1
    assert rows[0]["full_name"] == "Current name"
    assert rows[0]["event_count"] == 1
