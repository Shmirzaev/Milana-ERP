from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import Employee
from app.tests.conftest import TestSessionLocal
from app.tests.test_payroll import _create_user_with_permissions


@pytest.mark.parametrize("lang,title", [("en", "Employees"), ("ru", "Сотрудники"), ("uz", "Xodimlar")])
def test_export_all_matching_rows_literal_values_and_factory_scope(client, auth_headers, lang, title):
    with TestSessionLocal() as db:
        db.add_all([Employee(factory_code="MIL", full_name=f"Export_% {i:03d}", employee_no=f"00{i:04d}", phone="+998123", position="=1+1") for i in range(121)])
        db.add(Employee(factory_code="BST", full_name="Export_% other factory"))
        db.add(Employee(factory_code="MIL", full_name="Export other search"))
        db.commit()
    result = client.get('/api/employees/export.xlsx', headers=auth_headers, params={"search": "Export_%", "lang": lang})
    assert result.status_code == 200, result.text
    assert result.headers['cache-control'] == 'no-store'
    book = load_workbook(BytesIO(result.content))
    sheet = book[title]
    assert sheet.max_row == 122 and sheet.max_column == 9
    assert sheet["A2"].value == "000120"
    assert sheet["E2"].value == "=1+1" and sheet["E2"].data_type == 's'
    assert sheet["I2"].value == "+998123" and sheet["I2"].data_type == 's'
    assert sheet.freeze_panes == 'C2' and sheet.auto_filter.ref == 'A1:I122'
    listed = client.get('/api/employees', headers=auth_headers, params={"page": 1, "search": "Export_%"}).json()
    assert listed['total'] == sheet.max_row - 1


def test_export_matches_list_private_visibility_and_authentication(client, auth_headers):
    with TestSessionLocal() as db:
        db.add(Employee(factory_code="MIL", full_name="Export privacy", phone="private-phone", salary=12345, hr_profile_json={"bank_details": "private-bank"}))
        db.commit()
    viewer = _create_user_with_permissions(client, auth_headers, email='export-viewer@example.com', permissions=['payroll.view'])
    manager = _create_user_with_permissions(client, auth_headers, email='export-manager@example.com', permissions=['payroll.manage'])
    for headers, columns in [(viewer, 8), (manager, 9)]:
        result = client.get('/api/employees/export.xlsx?search=Export%20privacy', headers=headers)
        assert result.status_code == 200, result.text
        sheet = load_workbook(BytesIO(result.content)).active
        assert sheet.max_column == columns and sheet.max_row == 2
        values = str(list(sheet.values))
        assert 'private-bank' not in values and '12345' not in values
        assert ('private-phone' in values) == (columns == 9)
    client.cookies.clear()
    assert client.get('/api/employees/export.xlsx').status_code == 401


def test_empty_export_retains_headers(client, auth_headers):
    result = client.get('/api/employees/export.xlsx?search=nonexistent-export-person', headers=auth_headers)
    assert load_workbook(BytesIO(result.content)).active.max_row == 1
