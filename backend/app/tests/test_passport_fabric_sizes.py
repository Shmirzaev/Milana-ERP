from io import BytesIO

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.schemas import Bundle, CuttingPassport, StockBatch
from app.tests.test_passport_cutting_submission import passport_cutting  # noqa: F401


def test_independent_sizes_round_trip_export_and_cutting_handoff(client, auth_headers, passport_cutting):
    batches, order, _, cutting = passport_cutting
    url = f"/api/cutting-passports/{cutting['cutting_passport_id']}"
    saved = client.get(url, headers=auth_headers).json()
    form = {**saved, "size_range": "46", "materials": [
        {**row, "size_range": selection, "fabric_width_m": 2, "lay_length_m": 4, "gramage": 0.2}
        for row, selection in zip(saved["materials"], ["46", "46 ×2"])
    ]}
    result = client.patch(url, headers=auth_headers, json=form)
    assert result.status_code == 200, result.text
    for data in [result.json(), client.get(url, headers=auth_headers).json()]:
        assert [r["size_range"] for r in data["materials"]] == ["46", "46 ×2"]
        assert [r["size_count"] for r in data["materials"]] == [1, 2]
        assert [r["per_piece_weight_kg"] for r in data["materials"]] == [1.6, 0.8]
        assert data["size_range"] == "46" and data["size_count"] == 1
    exported = client.get('/api/cutting-passports/export.xlsx', params={"q": saved["passport_no"]}, headers=auth_headers)
    book = load_workbook(BytesIO(exported.content))
    assert [book["Materials"].cell(i, 7).value for i in (2, 3)] == ["46", "46 ×2"]
    # Editing the second fabric must leave the first fabric untouched.
    form["materials"][1]["size_range"] = "46"
    assert client.patch(url, headers=auth_headers, json=form).json()["materials"][0]["size_count"] == 1
    form["materials"][1]["size_range"] = "46 ×2"
    assert client.patch(url, headers=auth_headers, json=form).status_code == 200
    result = client.post('/api/cutting/records', headers=auth_headers, json=cutting)
    assert result.status_code == 201, result.text
    assert [r["quantity"] for r in result.json()["materials"]] == [6.5, 6.5]
    with SessionLocal() as db:
        assert [float(db.get(StockBatch, b["id"]).quantity) for b in batches] == [93.5, 93.5]
        assert db.query(Bundle).filter_by(production_order_id=order["id"]).one().quantity == 10


def test_legacy_rows_use_passport_size_without_rewriting_saved_data(client, auth_headers, passport_cutting):
    _, _, _, cutting = passport_cutting
    with SessionLocal() as db:
        passport = db.get(CuttingPassport, cutting["cutting_passport_id"])
        passport.size_range = "46 ×2"
        passport.materials = [{k: v for k, v in row.items() if k != "size_range"} for row in passport.materials]
        db.commit()
    response = client.get(f"/api/cutting-passports/{cutting['cutting_passport_id']}", headers=auth_headers)
    assert [r["size_count"] for r in response.json()["materials"]] == [2, 2]
    assert [r["size_range"] for r in response.json()["materials"]] == ["46 ×2", "46 ×2"]
    with SessionLocal() as db:
        assert all("size_range" not in row for row in db.get(CuttingPassport, cutting["cutting_passport_id"]).materials)
