from datetime import datetime, timezone
from io import BytesIO
import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import event
from sqlalchemy import create_engine, text
from alembic.migration import MigrationContext
from alembic.operations import Operations

from app.db.session import SessionLocal
from app.models import Bundle, CuttingPassport, ProductionBatch, StockBatch, WorkOrder
from app.schemas.cutting_material import PassportMaterial
from app.tests.test_passport_cutting_submission import passport_cutting  # noqa: F401
from app.tests.test_passport_get_scope import _actor, _headers, _passport, _snapshot


def _create(client, headers, **changes):
    payload = {"passport_no": "REPORT", "date": "2026-10-17T00:00:00Z", "pieces": 300,
               "layer_weight_kg": 2, "total_layers": 5, "rolls_count": 3, "beka_per_piece_kg": 0.01, "scrap_kg": 1}
    response = client.post("/api/cutting-passports", headers=headers, json=payload | changes)
    assert response.status_code == 201, response.text
    return response.json()


def test_tepa_pass_roundtrip_dashboard_and_excel_keep_original_quantities(client, auth_headers):
    saved = _create(client, auth_headers, tepa_pass=True, notes="=HYPERLINK(\"https://invalid\")")
    assert saved["tepa_pass"] is True and saved["pieces"] == 300 and saved["actual_kg"] == 14
    report = client.get("/api/cutting-passports/dashboard", headers=auth_headers)
    assert report.status_code == 200, report.text
    assert report.json()["totals"] == dict(passports=1, pieces=600, fabric_kg=14, fabric_m=0, rolls=3)
    result = client.get("/api/cutting-passports/export.xlsx", headers=auth_headers)
    assert result.status_code == 200, result.text
    book = load_workbook(BytesIO(result.content))
    assert book["Passports"]["K2"].value == 300
    assert book["Passports"]["J2"].value is True
    assert book["Passports"]["L2"].value.startswith("=HYPERLINK")
    assert book["Passports"]["L2"].data_type == "s"
    assert book["Materials"]["R2"].value == 300
    saved["tepa_pass"] = False
    updated = client.patch(f"/api/cutting-passports/{saved['id']}", headers=auth_headers, json=saved)
    assert updated.status_code == 200, updated.text
    assert updated.json()["pieces"] == 300 and updated.json()["actual_kg"] == 14
    assert client.get("/api/cutting-passports/dashboard", headers=auth_headers).json()["totals"]["pieces"] == 300


def test_mixed_fabrics_count_pieces_once_and_do_not_mix_kg_with_metres(client, auth_headers):
    saved = _create(client, auth_headers, tepa_pass=True)
    with SessionLocal() as db:
        p = db.get(CuttingPassport, saved["id"])
        p.materials = [PassportMaterial(stock_batch_id=index, meter_mode=metres, pieces=300,
                                      layer_weight_kg=weight, total_layers=5, rolls_count=rolls).model_dump()
                       for index, metres, weight, rolls in [(1, False, 2, 3), (2, True, 8, 4), (3, False, 1, 2)]]
        db.commit()
    data = client.get("/api/cutting-passports/dashboard", headers=auth_headers).json()
    assert data["totals"] == dict(passports=1, pieces=600, fabric_kg=15, fabric_m=40, rolls=9)
    assert data["daily"] == [{"date": "2026-10-17", **data["totals"]}]
    book = load_workbook(BytesIO(client.get("/api/cutting-passports/export.xlsx", headers=auth_headers).content))
    assert book["Passports"].max_row == 2 and book["Materials"].max_row == 4
    assert [book["Materials"].cell(row, 11).value for row in [2, 3, 4]] == ["kg", "m", "kg"]


def test_export_and_dashboard_include_more_than_500_and_date_end_is_inclusive(client, auth_headers):
    with SessionLocal() as db:
        db.add_all([CuttingPassport(passport_no=f"MANY-{i}", date=datetime(2026, 10, 17, 23, 59, 59, tzinfo=timezone.utc),
                                   pieces=1, rolls_count=1, layer_weight_kg=0.1, total_layers=1) for i in range(503)])
        db.add(CuttingPassport(passport_no="OTHER-DATE", date=datetime(2026, 10, 18, tzinfo=timezone.utc), pieces=1000))
        db.commit()
    query = "?date_from=2026-10-17&date_to=2026-10-17"
    assert len(client.get("/api/cutting-passports" + query + "&limit=500", headers=auth_headers).json()) == 500
    dashboard = client.get("/api/cutting-passports/dashboard" + query, headers=auth_headers)
    assert dashboard.status_code == 200, dashboard.text
    assert dashboard.json()["totals"] == dict(passports=503, pieces=503, fabric_kg=50.3, fabric_m=0, rolls=503)
    export = client.get("/api/cutting-passports/export.xlsx" + query, headers=auth_headers)
    assert load_workbook(BytesIO(export.content))["Passports"].max_row == 504
    assert client.get("/api/cutting-passports/dashboard?q=OTHER", headers=auth_headers).json()["totals"]["pieces"] == 1000


@pytest.mark.parametrize("endpoint", ["", "/dashboard", "/export.xlsx"])
def test_bad_dates_are_rejected(client, auth_headers, endpoint):
    url = "/api/cutting-passports" + endpoint
    assert client.get(url + "?date_from=2026-10-18&date_to=2026-10-17", headers=auth_headers).status_code == 422
    assert client.get(url + "?date_from=invalid", headers=auth_headers).status_code == 422


@pytest.mark.parametrize("endpoint", ["/dashboard", "/export.xlsx"])
def test_reports_enforce_permission_and_factory_and_are_read_only(client, endpoint):
    mil = _passport("MIL")
    eco = _passport("ECO")
    before = [_snapshot(mil), _snapshot(eco)]
    assert client.get("/api/cutting-passports" + endpoint).status_code == 401
    assert client.get("/api/cutting-passports" + endpoint, headers=_headers(_actor(permissions=()), "MIL")).status_code == 403
    actor = _actor("MIL")
    headers = _headers(actor, "MIL")
    assert client.get("/api/cutting-passports" + endpoint + "?cutting_department_code=ECT", headers=headers).status_code == 403
    writes = []
    with SessionLocal() as db:
        engine = db.get_bind()
    def track(connection, cursor, statement, parameters, context, executemany):
        if statement.lstrip().split(None, 1)[0].upper() in {"INSERT", "UPDATE", "DELETE"}:
            writes.append(statement)
    event.listen(engine, "before_cursor_execute", track)
    try:
        response = client.get("/api/cutting-passports" + endpoint, headers=headers)
    finally:
        event.remove(engine, "before_cursor_execute", track)
    assert response.status_code == 200, response.text[:300]
    if endpoint == "/dashboard":
        assert response.json()["totals"]["passports"] == 1
    else:
        sheet = load_workbook(BytesIO(response.content))["Passports"]
        assert sheet.max_row == 2 and sheet["L2"].value == "MIL private cutting notes"
    assert writes == [] and [_snapshot(mil), _snapshot(eco)] == before
    eco_response = client.get("/api/cutting-passports" + endpoint, headers=_headers(_actor("ECO"), "ECO"))
    assert eco_response.status_code == 200
    if endpoint == "/dashboard":
        assert eco_response.json()["totals"]["passports"] == 1
    else:
        assert load_workbook(BytesIO(eco_response.content))["Passports"]["L2"].value == "ECO private cutting notes"


def test_tepa_pass_does_not_change_stock_nastil_bundles_or_downstream_plan(client, auth_headers, passport_cutting):
    batches, order, work, cutting = passport_cutting
    pid = cutting["cutting_passport_id"]
    saved = client.get(f"/api/cutting-passports/{pid}", headers=auth_headers).json()
    saved["tepa_pass"] = True
    result = client.patch(f"/api/cutting-passports/{pid}", headers=auth_headers, json=saved)
    assert result.status_code == 200, result.text
    assert result.json()["pieces"] == 10
    result = client.post("/api/cutting/records", headers=auth_headers, json=cutting)
    assert result.status_code == 201, result.text
    with SessionLocal() as db:
        assert [float(db.get(StockBatch, row["id"]).quantity) for row in batches] == [93.5, 93.5]
        assert db.query(Bundle).filter_by(production_order_id=order["id"]).one().quantity == 10
        assert db.get(WorkOrder, work["id"]).passed_qty == 10
        assert db.get(ProductionBatch, cutting["production_batch_id"]).planned_quantity == 10
        assert all(row.planned_output_qty == 10 for row in db.query(WorkOrder).filter_by(production_order_id=order["id"]))
    totals = client.get("/api/cutting-passports/dashboard", headers=auth_headers).json()["totals"]
    assert totals["pieces"] == 20 and totals["fabric_kg"] == 13 and totals["rolls"] == 2


@pytest.mark.parametrize("lang", ["ru", "uz"])
def test_export_localized_and_empty(client, auth_headers, lang):
    response = client.get(f"/api/cutting-passports/export.xlsx?lang={lang}", headers=auth_headers)
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.worksheets[0].max_row == 1
    assert book.worksheets[0]["B1"].value == {"ru": "Дата", "uz": "Sana"}[lang]
    totals = client.get("/api/cutting-passports/dashboard", headers=auth_headers).json()["totals"]
    assert all(value == 0 for value in totals.values())


def test_migration_preserves_existing_quantities_and_defaults_false():
    path = Path(__file__).parents[2] / "alembic/versions/0147_passport_tepa_pass.py"
    spec = importlib.util.spec_from_file_location("passport_tepa_migration", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE cutting_passports (id INTEGER PRIMARY KEY, pieces INTEGER)"))
        connection.execute(text("INSERT INTO cutting_passports VALUES (1, 300)"))
        migration.op = Operations(MigrationContext.configure(connection))
        migration.upgrade()
        assert connection.execute(text("SELECT pieces, tepa_pass FROM cutting_passports")).one() == (300, 0)
        connection.execute(text("INSERT INTO cutting_passports (id, pieces) VALUES (2, 500)"))
        assert connection.execute(text("SELECT pieces, tepa_pass FROM cutting_passports WHERE id=2")).one() == (500, 0)
    engine.dispose()
