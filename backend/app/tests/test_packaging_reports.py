from datetime import datetime, timezone
from io import BytesIO
import base64

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, func

from app.core.security import create_access_token
from app.models import (
    AuditLog,
    Brand,
    CuttingPassport,
    Department,
    LegacyStockReceipt,
    Model,
    ModelImage,
    Package,
    PackageBatchAllocation,
    PackageItem,
    ProductionBatch,
    ProductionOrder,
    Role,
    User,
    WorkOrder,
)
from app.tests.conftest import TestSessionLocal, test_engine


PARAMS = {"from_date": "2026-09-03", "to_date": "2026-09-03", "packaging_department_code": "PKG"}


@pytest.fixture
def report_data():
    with TestSessionLocal() as db:
        brand = Brand(name="Report brand")
        model = Model(
            code="REPORT-123",
            name="Report garment",
            category="Tunic",
            details_json={"general": {"model_no": "REPORT", "variant_no": "V-123"}},
        )
        db.add_all([brand, model])
        db.flush()
        db.add(
            ModelImage(
                model_id=model.id,
                file_url="/storage/model-files/report.png",
                file_name="report.png",
                content_type="image/png",
                image_type="model",
                is_primary=True,
                file_data=base64.b64decode(
                    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
                ),
            )
        )
        order = ProductionOrder(
            production_no="PO-REPORT",
            production_type="branded_stock",
            model_id=model.id,
            brand_id=brand.id,
            planned_quantity=180,
        )
        db.add(order)
        db.flush()
        batches = [
            ProductionBatch(production_order_id=order.id, batch_no=f"REPORT-{n}", batch_index=n, planned_quantity=90)
            for n in [1, 2]
        ]
        db.add_all(batches)
        db.flush()
        department = db.query(Department).filter_by(code="PKG").one()
        work = WorkOrder(
            production_order_id=order.id,
            department_id=department.id,
            operation="packaging",
            status="completed",
            planned_output_qty=180,
            passed_qty=175,
            failed_qty=3,
            end_time=datetime(2026, 9, 3, 18, 59, 59, tzinfo=timezone.utc),
            notes="=Do not execute",
        )
        db.add(work)
        db.add_all(
            [
                CuttingPassport(passport_no=n, date=datetime(2026, 9, 1), production_order_id=order.id, pieces=90)
                for n in ["P10", "P2"]
            ]
        )
        # Local September 3 is UTC September 2 19:00 through September 3 19:00.
        specs = [
            ("before", datetime(2026, 9, 2, 18, 59, 59), 90, "PKG"),
            ("start", datetime(2026, 9, 2, 19), 60, "PKG"),
            ("partial", datetime(2026, 9, 3, 12), 17, "PKG"),
            ("merged", datetime(2026, 9, 3, 18, 59, 59), 60, "PKG"),
            ("end", datetime(2026, 9, 3, 19), 90, "PKG"),
            ("foreign", datetime(2026, 9, 3, 12), 90, "ECP"),
        ]
        for name, packed_at, quantity, scope in specs:
            package = Package(
                package_no=f"REPORT-{name}",
                barcode=f"REPORT-{name}",
                production_order_id=order.id,
                production_batch_id=batches[0].id if name != "merged" else None,
                model_id=model.id,
                color="Blue",
                total_quantity=quantity,
                capacity=90,
                packaging_department_code=scope,
                packed_at=packed_at,
                received_at=datetime(2026, 9, 5, 12),
                status="shipped",
            )
            db.add(package)
            db.flush()
            db.add(PackageItem(package_id=package.id, model_id=model.id, color="Blue", size="M-46", quantity=quantity))
            if name == "merged":
                db.add_all(
                    [
                        PackageBatchAllocation(package_id=package.id, production_batch_id=b.id, quantity=30)
                        for b in batches
                    ]
                )
        receipt = LegacyStockReceipt(
            source_system="test",
            source_warehouse_id="test",
            source_record_id="old",
            source_checksum="test",
            source_payload={},
        )
        db.add(receipt)
        db.flush()
        db.add(
            Package(
                package_no="REPORT-import",
                barcode="REPORT-import",
                legacy_receipt_id=receipt.id,
                model_id=model.id,
                color="Blue",
                total_quantity=999,
                capacity=999,
                packaging_department_code="PKG",
                packed_at=datetime(2026, 9, 3, 12),
            )
        )
        db.commit()
        return order.id


def test_report_counts_actual_packages_once_and_uses_tashkent_creation_date(client, auth_headers, report_data):
    with TestSessionLocal() as db:
        before = (db.query(func.count(Package.id)).scalar(), db.query(func.count(AuditLog.id)).scalar())
    selects = []

    def count_select(_conn, _cursor, statement, _parameters, _context, _executemany):
        if statement.lstrip().upper().startswith("SELECT"):
            selects.append(statement)

    event.listen(test_engine, "before_cursor_execute", count_select)
    try:
        response = client.get("/api/packaging/reports", params=PARAMS, headers=auth_headers)
    finally:
        event.remove(test_engine, "before_cursor_execute", count_select)
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["totals"] == {"package_count": 3, "quantity": 137}
    assert data["daily"][0]["date"] == "2026-09-03"
    assert data["daily"][0]["quantity"] == 137
    assert data["daily"][0]["two_piece_quantity"] is None
    assert len(data["entries"]) == 2
    assert all(row["passport"] == "P2, P10" for row in data["entries"])
    assert all(row["variant_no"] == "V-123" and row["standards"] == "90" for row in data["entries"])
    merged = next(row for row in data["entries"] if row["package_count"] == 1)
    assert merged["batch"] == "REPORT-1, REPORT-2" and merged["quantity"] == 60
    closed = data["completed"][0]
    assert closed["packed_quantity"] == 175 and closed["damaged_quantity"] == 3 and closed["balance"] == -2
    assert closed["shortage"] is None and closed["first_sort"] is None
    assert len(selects) <= 13
    with TestSessionLocal() as db:
        after = (db.query(func.count(Package.id)).scalar(), db.query(func.count(AuditLog.id)).scalar())
        assert before == after


@pytest.mark.parametrize("lang", ["en", "ru", "uz"])
def test_excel_matches_report_and_preserves_unknowns_and_literal_text(client, auth_headers, report_data, lang):
    from app.services.packaging_report_exports import TEXT

    response = client.get("/api/packaging/reports/export.xlsx", params={**PARAMS, "lang": lang}, headers=auth_headers)
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == [TEXT[lang][key] for key in ["entries", "completed", "daily"]]
    assert len(workbook.worksheets[0]._images) == 2
    daily = workbook.worksheets[2]
    assert daily["B4"].value == datetime(2026, 9, 3)
    assert daily["C4"].value == 3 and daily["D4"].value == 137
    assert daily["E4"].value is None and daily["F4"].value is None
    assert daily["D5"].value == "=SUM(D4:D4)"
    closed = workbook.worksheets[1]
    assert closed["Q4"].value == "=Do not execute" and closed["Q4"].data_type == "s"
    assert daily.freeze_panes == "C4" and daily.auto_filter.ref == "A3:H4"


@pytest.mark.parametrize("endpoint", ["", "/export.xlsx"])
def test_report_requires_permission_and_selected_factory(client, report_data, endpoint):
    with TestSessionLocal() as db:
        role = Role(name="Report test", permissions=["packaging.records"])
        db.add(role)
        db.flush()
        user = User(
            name="Report test",
            email="report@example.com",
            password_hash="unused",
            role_id=role.id,
            department_id=db.query(Department.id).filter_by(code="PKG").scalar(),
            factory_code="MIL",
        )
        db.add(user)
        db.commit()
        headers = {"Authorization": f"Bearer {create_access_token(user.id, extra={'factory_code': 'MIL'})}"}
        assert client.get(f"/api/packaging/reports{endpoint}", params=PARAMS, headers=headers).status_code == 200
        for department in ["ECP", "BPK"]:
            assert (
                client.get(
                    f"/api/packaging/reports{endpoint}",
                    params={**PARAMS, "packaging_department_code": department},
                    headers=headers,
                ).status_code
                == 403
            )
        role.permissions = ["sales.orders"]
        db.commit()
        assert client.get(f"/api/packaging/reports{endpoint}", params=PARAMS, headers=headers).status_code == 403
    assert client.get(f"/api/packaging/reports{endpoint}", params=PARAMS).status_code == 401


@pytest.mark.parametrize("department,factory,qty", [("PKG", "MIL", 137), ("ECP", "ECO", 90), ("BPK", "BST", 0)])
def test_each_factory_report_is_isolated(client, report_data, department, factory, qty):
    headers = {"Authorization": f"Bearer {create_access_token(1, extra={'factory_code': factory})}"}
    response = client.get(
        "/api/packaging/reports", params={**PARAMS, "packaging_department_code": department}, headers=headers
    )
    assert response.status_code == 200, response.text
    assert response.json()["totals"]["quantity"] == qty


def test_empty_invalid_and_oversized_reports(client, auth_headers, report_data, monkeypatch):
    assert (
        client.get(
            "/api/packaging/reports", params={**PARAMS, "from_date": "2026-09-04"}, headers=auth_headers
        ).status_code
        == 400
    )
    assert (
        client.get(
            "/api/packaging/reports", params={**PARAMS, "from_date": "invalid"}, headers=auth_headers
        ).status_code
        == 422
    )
    empty = client.get(
        "/api/packaging/reports",
        params={**PARAMS, "from_date": "2026-09-06", "to_date": "2026-09-06"},
        headers=auth_headers,
    )
    assert empty.status_code == 200 and empty.json()["daily"] == []
    monkeypatch.setattr("app.services.packaging_reports.MAX_REPORT_PACKAGES", 1)
    assert client.get("/api/packaging/reports", params=PARAMS, headers=auth_headers).status_code == 413
