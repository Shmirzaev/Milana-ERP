from io import BytesIO

from openpyxl import load_workbook


def test_qolip_filter_keeps_family_and_matches_legacy_fields(client, auth_headers):
    from app.db import session as session_module
    from app.models import Model

    with session_module.SessionLocal() as db:
        for number, variant, fields in [
            ("QF991", "1", {"qolip_no": "PAT%_991"}),
            ("QF991", "2", {}),
            ("QF992", "1", {"moldNo": "pat%_991"}),
            ("QF993", "1", {"qolip_no": "PATxxx991"}),
            ("QF994", "1", {"qolip_no": "current", "mold_no": "PAT%_991"}),
        ]:
            db.add(Model(code=f"{number}-{variant}", name="Qolip fixture", status="approved",
                         details_json={"general": {"model_no": number, "variant_no": variant, **fields}}, created_by=1))
        db.commit()
    response = client.get("/api/models/variant-groups", headers=auth_headers,
                          params={"qolip_no": " pat%_991 ", "include_total": True, "compact": True})
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["total"] == 2
    assert {row["group_model_no"]: row["variant_count"] for row in payload["rows"]} == {"QF991": 2, "QF992": 1}


def export_payload():
    return {"model": "PJ001", "order": "PO-001", "factory": "MIL", "currency": "UZS",
            "headers": ["No.", "Use", "Section", "Code", "Name", "Rate", "Copies", "Divide"],
            "rows": [{"selected": True, "section": "Sewing", "code": "0012", "name": '=HYPERLINK("bad")',
                      "rate": "123.4567", "copies": 2, "division": "No divide"},
                     {"selected": False, "section": "Cutting", "code": "0013", "name": "Cut\u0001 cloth",
                      "rate": "0", "copies": 1, "division": "Equal"}]}


def test_process_export_is_real_excel_with_literal_text_and_numeric_rates(client, auth_headers):
    response = client.post("/api/payroll/process-qr/export.xlsx", headers=auth_headers, json=export_payload())
    assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.max_row == 4
    assert sheet["D3"].value == "0012"
    assert sheet["E3"].data_type == "s"
    assert sheet["F3"].value == 123.4567
    assert sheet["B4"].value is False
    assert sheet["E4"].value == "Cut cloth"


def test_process_export_rejects_other_factory(client, auth_headers):
    payload = export_payload()
    payload["factory"] = "BST"
    response = client.post("/api/payroll/process-qr/export.xlsx", headers=auth_headers, json=payload)
    assert response.status_code == 403
