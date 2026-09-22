from datetime import date, timedelta
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import event

from app.core.security import create_access_token
from app.db.session import SessionLocal
from app.models import AuditLog, Role, SewingDailyReport, SewingFlow, User
from app.tests.conftest import test_engine


LATEST_DAY = date(2099, 12, 31)


def _seed_reports(count: int) -> tuple[date, list[str]]:
    marker = uuid4().hex[:10].upper()
    with SessionLocal() as db:
        db.query(SewingDailyReport).delete(synchronize_session=False)
        flow = SewingFlow(
            factory_code="MIL",
            name=f"Export line {marker}",
            code=f"X-{marker}",
            capacity_per_day=500,
            is_active=True,
        )
        db.add(flow)
        db.flush()
        model_numbers = [f"EXPORT-{marker}-{index:04d}" for index in range(count)]
        db.add_all([
            SewingDailyReport(
                report_date=LATEST_DAY - timedelta(days=index),
                sewing_flow_id=flow.id,
                line_code=flow.code,
                line_name=flow.name,
                manual_model_no=model_numbers[index],
                manual_variant_no=f"V-{index:04d}",
                kroy_no=f"KR-{marker}-{index:04d}",
                sewn_qty=index + 1,
                defective_qty=0,
            )
            for index in range(count)
        ])
        db.commit()
        return LATEST_DAY - timedelta(days=count - 1), model_numbers


def _entry_rows(content: bytes) -> list[list[object]]:
    sheet = load_workbook(BytesIO(content), read_only=True, data_only=False)["Entries"]
    return [
        [sheet.cell(row, column).value for column in range(1, 6)]
        for row in range(6, sheet.max_row + 1)
    ]


def _capture_export(client, auth_headers, *, start: date, page: int, page_size: int):
    statements: list[str] = []

    def capture(_conn, _cursor, statement, _parameters, _context, _executemany):
        normalized = " ".join(statement.lower().split())
        if normalized.startswith("select"):
            statements.append(normalized)

    event.listen(test_engine, "before_cursor_execute", capture)
    try:
        response = client.get(
            "/api/sewing-daily-reports/export.xlsx",
            params={
                "from_date": start.isoformat(),
                "to_date": LATEST_DAY.isoformat(),
                "lang": "en",
                "page": page,
                "page_size": page_size,
            },
            headers=auth_headers,
        )
    finally:
        event.remove(test_engine, "before_cursor_execute", capture)
    return response, statements


def test_sewing_daily_xlsx_page_is_sql_bounded_at_1_50_and_401_rows(client, auth_headers):
    query_counts: list[int] = []
    for count in (1, 50, 401):
        start, model_numbers = _seed_reports(count)
        response, statements = _capture_export(
            client, auth_headers, start=start, page=1, page_size=25,
        )

        assert response.status_code == 200, response.text
        rows = _entry_rows(response.content)
        assert [row[2] for row in rows] == model_numbers[:25]
        assert [row[0] for row in rows] == list(range(1, min(count, 25) + 1))
        assert response.headers["X-Total-Count"] == str(count)
        assert response.headers["X-Page"] == "1"
        assert response.headers["X-Page-Size"] == "25"
        assert response.headers["X-Has-More"] == ("true" if count > 25 else "false")
        report_reads = [statement for statement in statements if "sewing_daily_reports" in statement]
        assert any(" limit ? offset ?" in statement for statement in report_reads), report_reads
        query_counts.append(len(statements))

    print(f"Sewing daily XLSX SELECTs at 1/50/401 rows: {query_counts}")
    assert max(query_counts) - min(query_counts) <= 1, query_counts


def test_sewing_daily_xlsx_page_matches_legacy_order_and_rows(client, auth_headers):
    start, model_numbers = _seed_reports(55)
    common = {
        "from_date": start.isoformat(),
        "to_date": LATEST_DAY.isoformat(),
        "lang": "en",
    }
    legacy = client.get(
        "/api/sewing-daily-reports/export.xlsx", params=common, headers=auth_headers,
    )
    paged = client.get(
        "/api/sewing-daily-reports/export.xlsx",
        params={**common, "page": 2, "page_size": 20},
        headers=auth_headers,
    )

    assert legacy.status_code == 200, legacy.text
    assert paged.status_code == 200, paged.text
    assert "X-Total-Count" not in legacy.headers
    legacy_rows = _entry_rows(legacy.content)
    page_rows = _entry_rows(paged.content)
    assert [row[2] for row in legacy_rows] == model_numbers
    assert [row[1:] for row in page_rows] == [row[1:] for row in legacy_rows[20:40]]
    assert [row[0] for row in page_rows] == list(range(1, 21))
    assert paged.headers["X-Total-Count"] == "55"


def test_sewing_daily_xlsx_page_auth_validation_and_no_writes(client, auth_headers):
    start, _model_numbers = _seed_reports(3)
    with SessionLocal() as db:
        role = Role(name=f"No sewing report export {uuid4().hex}", permissions=[])
        db.add(role)
        db.flush()
        user = User(
            name="Denied sewing report export",
            email=f"denied-sewing-report-{uuid4().hex}@example.com",
            password_hash="unused-pagination-hash",
            role_id=role.id,
            factory_code="MIL",
            extra_permissions=[],
            is_active=True,
        )
        db.add(user)
        db.commit()
        denied_headers = {"Authorization": f"Bearer {create_access_token(user.id)}"}
        before = (db.query(SewingDailyReport).count(), db.query(AuditLog).count())

    params = {
        "from_date": start.isoformat(),
        "to_date": LATEST_DAY.isoformat(),
        "page": 1,
        "page_size": 2,
    }
    assert client.get("/api/sewing-daily-reports/export.xlsx", params=params).status_code == 401
    assert client.get(
        "/api/sewing-daily-reports/export.xlsx", params=params, headers=denied_headers,
    ).status_code == 403
    assert client.get(
        "/api/sewing-daily-reports/export.xlsx",
        params={**params, "page": 0},
        headers=auth_headers,
    ).status_code == 422
    assert client.get(
        "/api/sewing-daily-reports/export.xlsx",
        params={**params, "page_size": 501},
        headers=auth_headers,
    ).status_code == 422
    allowed = client.get(
        "/api/sewing-daily-reports/export.xlsx", params=params, headers=auth_headers,
    )
    assert allowed.status_code == 200, allowed.text
    assert allowed.headers["X-Total-Count"] == "3"

    with SessionLocal() as db:
        after = (db.query(SewingDailyReport).count(), db.query(AuditLog).count())
    assert after == before
