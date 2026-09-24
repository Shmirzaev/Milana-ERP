from copy import deepcopy
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.api.routes.shipments import _invoice_print_details
from app.db.session import SessionLocal
from app.models import LegacyStockReceipt, Package, Shipment, ShipmentScanLog
from app.services.shipment_invoice import build_invoice_rows, render_shipment_invoice
from app.services.shipment_invoice_excel import shipment_invoice_workbook
from app.tests.test_shipment_review import dispatch as dispatch, ship


def test_partial_weight_excel_shared_pack_and_literal_text():
    packages = [{"package_no": "A", "quantity": 7, "weight_kg": "12.50"}, {"package_no": "B", "quantity": 3, "weight_kg": None}]
    lines = [{"package_no": "A", "model_no": "=1+1", "description": "First", "size": "48", "quantity": 4, "amount": None},
             {"package_no": "A", "model_no": "OTHER", "size": "50", "quantity": 3, "amount": None},
             {"package_no": "B", "model_no": "LAST", "size": "52", "quantity": 3, "amount": None}]
    document = {"shipment_no": "TEST", "package_details": packages, "lines": lines, "invoice_rows": build_invoice_rows(lines, packages),
                "packages_count": 2, "quantity": 10, "total_weight_kg": None, "known_weight_kg": "12.50", "missing_weight_packages": 1,
                "invoice_layout_version": 2, "finance_posting_status": "pending_price"}
    original = deepcopy(document)
    for lang in ("en", "ru", "uz"):
        html = render_shipment_invoice(document, lang)
        assert html.count('class="numeric">12.50</td>') == 4
        sheet = load_workbook(BytesIO(shipment_invoice_workbook(document, lang))).active
        assert sheet["B10"].value == "=1+1" and sheet["B10"].data_type == "s"
        assert sheet["F13"].value == 2 and sheet["G13"].value == 10
        assert sheet["H13"].value == sheet["I13"].value == 12.5
        assert sheet["H12"].value is None
        assert "H10:H11" in sheet.merged_cells
    assert document == original


def test_invoice_scan_order_ignores_detached_and_duplicate_scans(client, auth_headers, dispatch):
    with SessionLocal() as db:
        first = db.get(Package, dispatch["package"])
        receipt = LegacyStockReceipt(source_system="TEST", source_warehouse_id="scan-order", source_record_id="second",
                                     source_checksum="b" * 64, source_payload={"quantity": 1})
        db.add(receipt); db.flush()
        second = Package(package_no="SCAN-SECOND", barcode="SCAN-SECOND", model_id=first.model_id,
                         legacy_receipt_id=receipt.id, color="navy", capacity=60, total_quantity=1, status="packed")
        db.add(second); db.flush()
        start = datetime.now(timezone.utc)
        original_scan = db.query(ShipmentScanLog).filter_by(shipment_id=dispatch["shipment"]).one()
        original_scan.scanned_at = start
        for offset, package, result in [(1, second, "matched"), (2, first, "detached"), (3, first, "matched"), (4, second, "matched")]:
            db.add(ShipmentScanLog(shipment_id=dispatch["shipment"], package_id=package.id, scanned_code=package.barcode,
                                  scan_result=result, scanned_at=start + timedelta(seconds=offset)))
        db.flush()
        document = {"package_details": [{"package_no": first.package_no}, {"package_no": second.package_no}], "lines": []}
        result = _invoice_print_details(db, db.get(Shipment, dispatch["shipment"]), document)
        assert [p["package_no"] for p in result["package_details"]] == [second.package_no, first.package_no]


def test_excel_endpoint_reuses_frozen_invoice_and_does_not_write(client, auth_headers, dispatch):
    base = f'/api/shipments/{dispatch["shipment"]}'
    assert client.get(base + "/invoice.xlsx", headers=auth_headers).status_code == 409
    ship(client, auth_headers, dispatch)
    with SessionLocal() as db:
        before = deepcopy(db.get(Shipment, dispatch["shipment"]).dispatch_snapshot)
    assert client.get(base + "/invoice.xlsx").status_code == 401
    response = client.get(base + "/invoice.xlsx?lang=uz", headers=auth_headers)
    assert response.status_code == 200, response.text
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet["G11"].value == 8
    with SessionLocal() as db:
        assert db.get(Shipment, dispatch["shipment"]).dispatch_snapshot == before
