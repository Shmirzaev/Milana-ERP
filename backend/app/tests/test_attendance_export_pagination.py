from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import event

from app.core.security import create_access_token
from app.db.session import SessionLocal
from app.models import AttendanceDevice, AttendancePerson, AuditLog, Role, User
from app.tests.conftest import test_engine


DAY = "2026-08-17"


def _seed_people(count: int) -> list[str]:
    marker = uuid4().hex[:10].upper()
    with SessionLocal() as db:
        db.query(AttendancePerson).delete(synchronize_session=False)
        device = AttendanceDevice(
            factory_code="MIL",
            device_key=f"perf35-export-{marker}",
            name=f"Export device {marker}",
            vendor="Hikvision",
            read_only=True,
        )
        db.add(device)
        db.flush()
        external_ids = [f"{marker}-{index:04d}" for index in range(count)]
        db.add_all([
            AttendancePerson(
                factory_code="MIL",
                device_id=device.id,
                external_person_id=external_ids[index],
                full_name=f"Attendance export {index:04d}",
                present_on_device=True,
                last_synced_at=datetime(2026, 8, 17, tzinfo=timezone.utc),
            )
            for index in range(count)
        ])
        db.commit()
        return external_ids


def _workbook_rows(content: bytes) -> list[list[object]]:
    sheet = load_workbook(BytesIO(content), read_only=True, data_only=False).active
    return [
        [sheet.cell(row, column).value for column in range(1, 8)]
        for row in range(7, sheet.max_row + 1)
    ]


def _captured_export(client, auth_headers, *, page: int, page_size: int):
    statements: list[str] = []

    def capture(_conn, _cursor, statement, _parameters, _context, _executemany):
        normalized = " ".join(statement.lower().split())
        if normalized.startswith("select"):
            statements.append(normalized)

    event.listen(test_engine, "before_cursor_execute", capture)
    try:
        response = client.get(
            "/api/attendance/reports/daily.xlsx",
            params={"day": DAY, "lang": "en", "page": page, "page_size": page_size},
            headers=auth_headers,
        )
    finally:
        event.remove(test_engine, "before_cursor_execute", capture)
    return response, statements


def test_attendance_export_page_is_sql_bounded_at_1_50_and_401_rows(client, auth_headers):
    query_counts: list[int] = []
    for count in (1, 50, 401):
        external_ids = _seed_people(count)
        response, statements = _captured_export(client, auth_headers, page=1, page_size=25)

        assert response.status_code == 200, response.text
        rows = _workbook_rows(response.content)
        assert [row[1] for row in rows] == external_ids[:25]
        assert [row[0] for row in rows] == list(range(1, min(count, 25) + 1))
        assert response.headers["X-Total-Count"] == str(count)
        assert response.headers["X-Page"] == "1"
        assert response.headers["X-Page-Size"] == "25"
        assert response.headers["X-Has-More"] == ("true" if count > 25 else "false")
        attendance_reads = [statement for statement in statements if "attendance_people" in statement]
        assert any(" limit ? offset ?" in statement for statement in attendance_reads), attendance_reads
        query_counts.append(len(statements))

    print(f"Attendance XLSX export SELECTs at 1/50/401 rows: {query_counts}")
    assert max(query_counts) - min(query_counts) <= 1, query_counts


def test_attendance_export_page_matches_legacy_order_filters_and_rows(client, auth_headers):
    external_ids = _seed_people(55)
    legacy = client.get(
        "/api/attendance/reports/daily.xlsx",
        params={"day": DAY, "lang": "en"},
        headers=auth_headers,
    )
    paged = client.get(
        "/api/attendance/reports/daily.xlsx",
        params={"day": DAY, "lang": "en", "page": 2, "page_size": 20},
        headers=auth_headers,
    )

    assert legacy.status_code == 200, legacy.text
    assert paged.status_code == 200, paged.text
    assert "X-Total-Count" not in legacy.headers
    legacy_rows = _workbook_rows(legacy.content)
    page_rows = _workbook_rows(paged.content)
    assert [row[1] for row in legacy_rows] == external_ids
    assert [row[1:] for row in page_rows] == [row[1:] for row in legacy_rows[20:40]]
    assert [row[0] for row in page_rows] == list(range(1, 21))
    assert paged.headers["X-Total-Count"] == "55"


def test_attendance_export_page_auth_validation_and_no_writes(client, auth_headers):
    _seed_people(3)
    with SessionLocal() as db:
        denied_role = Role(name=f"No attendance export {uuid4().hex}", permissions=[])
        db.add(denied_role)
        db.flush()
        denied_user = User(
            name="Denied attendance export",
            email=f"denied-attendance-export-{uuid4().hex}@example.com",
            password_hash="unused-pagination-hash",
            role_id=denied_role.id,
            factory_code="MIL",
            extra_permissions=[],
            is_active=True,
        )
        db.add(denied_user)
        db.commit()
        denied_headers = {"Authorization": f"Bearer {create_access_token(denied_user.id)}"}
        before = (db.query(AttendancePerson).count(), db.query(AuditLog).count())

    params = {"day": DAY, "page": 1, "page_size": 2}
    assert client.get("/api/attendance/reports/daily.xlsx", params=params).status_code == 401
    assert client.get(
        "/api/attendance/reports/daily.xlsx", params=params, headers=denied_headers,
    ).status_code == 403
    assert client.get(
        "/api/attendance/reports/daily.xlsx",
        params={"day": DAY, "page": 0, "page_size": 2},
        headers=auth_headers,
    ).status_code == 422
    assert client.get(
        "/api/attendance/reports/daily.xlsx",
        params={"day": DAY, "page": 1, "page_size": 501},
        headers=auth_headers,
    ).status_code == 422
    allowed = client.get(
        "/api/attendance/reports/daily.xlsx", params=params, headers=auth_headers,
    )
    assert allowed.status_code == 200, allowed.text
    assert allowed.headers["X-Total-Count"] == "3"

    with SessionLocal() as db:
        after = (db.query(AttendancePerson).count(), db.query(AuditLog).count())
    assert after == before
