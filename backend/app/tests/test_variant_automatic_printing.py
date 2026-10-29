from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.schemas import ModelImage, ProductionOrder
from app.tests.conftest import TestSessionLocal
from app.tests.test_production_flow import _VALID_PNG


def attach_artwork(model_id=1, image_type="print"):
    with TestSessionLocal() as db:
        db.add(ModelImage(model_id=model_id, image_type=image_type, file_name="art.png",
                          file_url="/storage/model-files/art.png", content_type="image/png", file_data=_VALID_PNG))
        db.commit()


@pytest.mark.parametrize("url", ["/api/production-orders", "/api/planning/create-branded-production"])
@pytest.mark.parametrize("image_type,expected", [("print", True), ("material", False), ("model", False)])
def test_variant_artwork_creates_printing_without_manual_checkbox(client, auth_headers, url, image_type, expected):
    attach_artwork(image_type=image_type)
    response = client.post(url, headers=auth_headers, json={
        "production_type": "branded_stock", "model_id": 1, "planned_quantity": 8,
        "items": [{"model_id": 1, "color": "white", "size": "M", "planned_quantity": 8}],
    })
    assert response.status_code == 201, response.text
    result = client.get(f'/api/production-orders/{response.json()["id"]}', headers=auth_headers).json()
    assert result["items"][0]["printing_required"] is expected
    assert ("printing" in {w["operation"] for w in result["work_orders"]}) is expected
    assert sum(w["operation"] == "cutting" for w in result["work_orders"]) == 1


def test_bso_embeds_exact_variant_artwork(client, auth_headers):
    from app.models.branded_order_export_service import build_branded_order_workbook
    attach_artwork()
    response = client.post("/api/planning/create-branded-production", headers=auth_headers, json={
        "production_type": "branded_stock", "model_id": 1, "planned_quantity": 8,
        "items": [{"model_id": 1, "color": "white", "size": "M", "planned_quantity": 8}],
    })
    assert response.status_code == 201, response.text
    with TestSessionLocal() as db:
        po = db.get(ProductionOrder, response.json()["id"])
        book = load_workbook(BytesIO(build_branded_order_workbook(db, po.planning_order)))
    assert any(img.anchor._from.col == 9 for img in book.active._images)


def test_artwork_without_size_items_still_adds_printing_and_bso_flag(client, auth_headers):
    from app.models.branded_order_export_service import build_branded_order_workbook
    attach_artwork()
    response = client.post('/api/planning/create-branded-production', headers=auth_headers,
        json={"production_type": "branded_stock", "model_id": 1, "planned_quantity": 8})
    assert response.status_code == 201, response.text
    with TestSessionLocal() as db:
        po = db.get(ProductionOrder, response.json()['id'])
        assert any(w.operation == 'printing' for w in po.work_orders)
        book = load_workbook(BytesIO(build_branded_order_workbook(db, po.planning_order)))
        assert book.active['J3'].value == 'Ha'
        from PIL import Image
        embedded = next(img for img in book.active._images if img.anchor._from.col == 9)._data()
        assert Image.open(BytesIO(embedded)).convert('RGBA').tobytes() == Image.open(BytesIO(_VALID_PNG)).convert('RGBA').tobytes()


def test_client_sales_and_planning_inherit_variant_printing(client, auth_headers):
    attach_artwork()
    sale = client.post('/api/sales-orders', headers=auth_headers, json={"order_type": "client_order", "customer_id": 1,
        "items": [{"model_id": 1, "color": "white", "size": "M", "quantity": 8, "unit_price": 1}]})
    assert sale.status_code == 201, sale.text
    assert sale.json()['items'][0]['printing_required'] is True
    sid = sale.json()['id']
    assert client.post(f'/api/sales-orders/{sid}/confirm', headers=auth_headers).status_code == 200
    response = client.post('/api/planning/create-production-order', headers=auth_headers, json={
        "production_type": "client_order", "model_id": 1, "sales_order_id": sid, "planned_quantity": 8,
        "items": [{"model_id": 1, "color": "white", "size": "M", "planned_quantity": 8}]})
    assert response.status_code == 201, response.text
    detail = client.get(f'/api/production-orders/{response.json()["id"]}', headers=auth_headers).json()
    assert any(w['operation'] == 'printing' for w in detail['work_orders'])


def test_cutting_sheet_embeds_exact_variant_print(client, auth_headers):
    from app.tests.test_sewing_actual_output import make_batch
    from app.schemas import CuttingRecord
    poid, _, works = make_batch()
    with TestSessionLocal() as db:
        model_id = db.get(ProductionOrder, poid).model_id
        rid = db.query(CuttingRecord.id).filter_by(work_order_id=works['cutting']).scalar()
    attach_artwork(model_id=model_id)
    response = client.get(f'/api/cutting/records/{rid}/production-sheet', headers=auth_headers)
    assert response.status_code == 200, response.text
    import base64
    assert f"src='data:image/png;base64,{base64.b64encode(_VALID_PNG).decode()}' alt='Print artwork'" in response.text
