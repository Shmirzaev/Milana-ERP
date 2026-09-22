from datetime import datetime, timedelta, timezone
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import event

from app.db.session import SessionLocal
from app.models import AuditLog, Employee, PayrollRecord
from app.tests.conftest import test_engine


def _seed_records(count: int) -> tuple[int, list[str]]:
    marker = uuid4().hex[:10].upper()
    with SessionLocal() as db:
        employee = Employee(
            factory_code="MIL",
            employee_no=f"PERF35-{marker}",
            full_name=f"PERF35 export {marker}",
            position="Operator",
            status="active",
        )
        db.add(employee)
        db.flush()
        model_codes = [f"PERF35-{marker}-{index:04d}" for index in range(count)]
        start = datetime(2099, 12, 31, tzinfo=timezone.utc)
        db.add_all([
            PayrollRecord(
                factory_code="MIL",
                dedupe_key=uuid4().hex,
                employee_id=employee.id,
                model_code=model_codes[index],
                operation_section="sewing",
                operation_code="PERF35-OP",
                operation_name="PERF35 export",
                quantity=1,
                rate_per_piece=1,
                total_amount=1,
                currency="UZS",
                scanned_at=start + timedelta(seconds=index),
                status="recorded",
            )
            for index in range(count)
        ])
        db.commit()
        return employee.id, model_codes


def _export_models(content: bytes) -> list[str]:
    sheet = load_workbook(BytesIO(content), read_only=True, data_only=False)["Sewing report"]
    return [
        sheet.cell(row, 8).value
        for row in range(6, sheet.max_row - 1)
        if sheet.cell(row, 8).value
    ]


def _capture_export(client, auth_headers, employee_id: int, **params):
    statements: list[str] = []

    def capture(_conn, _cursor, statement, _parameters, _context, _executemany):
        normalized = " ".join(statement.lower().split())
        if normalized.startswith("select"):
            statements.append(normalized)

    event.listen(test_engine, "before_cursor_execute", capture)
    try:
        response = client.get(
            "/api/payroll/reports/sewing-production.xlsx",
            params={"employee_id": employee_id, **params},
            headers=auth_headers,
        )
    finally:
        event.remove(test_engine, "before_cursor_execute", capture)
    return response, statements


def test_payroll_sewing_xlsx_paging_matches_legacy_order_and_rows(client, auth_headers):
    employee_id, model_codes = _seed_records(55)
    legacy, _ = _capture_export(client, auth_headers, employee_id)
    paged, _ = _capture_export(client, auth_headers, employee_id, page=2, page_size=20)

    assert legacy.status_code == 200, legacy.text
    assert paged.status_code == 200, paged.text
    assert "X-Total-Count" not in legacy.headers
    assert _export_models(legacy.content) == list(reversed(model_codes))
    assert _export_models(paged.content) == list(reversed(model_codes))[20:40]
    assert paged.headers["X-Total-Count"] == "55"
    assert paged.headers["X-Page"] == "2"
    assert paged.headers["X-Page-Size"] == "20"
    assert paged.headers["X-Has-More"] == "true"


def test_payroll_sewing_xlsx_page_uses_sql_limit_and_keeps_auth_validation(client, auth_headers):
    query_counts = []
    for count in (1, 50, 401):
        employee_id, _ = _seed_records(count)
        response, statements = _capture_export(
            client, auth_headers, employee_id, page=1, page_size=25,
        )
        assert response.status_code == 200, response.text
        assert response.headers["X-Total-Count"] == str(count)
        report_reads = [statement for statement in statements if "from payroll_records" in statement]
        assert any(" limit ? offset ?" in statement for statement in report_reads), report_reads
        query_counts.append(len(statements))

    print(f"Payroll sewing XLSX SELECTs at 1/50/401 rows: {query_counts}")
    assert max(query_counts) - min(query_counts) <= 1, query_counts
    assert client.get("/api/payroll/reports/sewing-production.xlsx").status_code == 401
    assert client.get(
        "/api/payroll/reports/sewing-production.xlsx",
        params={"page": 0},
        headers=auth_headers,
    ).status_code == 422

    with SessionLocal() as db:
        before = (db.query(PayrollRecord).count(), db.query(AuditLog).count())
    allowed = client.get(
        "/api/payroll/reports/sewing-production.xlsx",
        params={"employee_id": employee_id, "page": 1, "page_size": 25},
        headers=auth_headers,
    )
    assert allowed.status_code == 200, allowed.text
    with SessionLocal() as db:
        after = (db.query(PayrollRecord).count(), db.query(AuditLog).count())
    assert after == before
