from datetime import datetime, timezone
from io import BytesIO
import pytest
from openpyxl import load_workbook
from app.db.session import SessionLocal
from app.schemas import Employee, Model, BrandedPlanningOrder, ProductionOrder, ProductionOrderItem, PayrollRecord, SewingFlow, WorkOrder, Department
from app.models.employee_names_service import latin_employee_name
from app.models.employee_reconciliation_service import reconcile_employees
from app.tests.test_payroll import _create_user_with_permissions


@pytest.mark.parametrize('source,expected', [('Алиева Муштарий Вохиджон қизи','Aliyeva Mushtariy Voxidjon qizi'),
    ('G`OFURJONOVA   MUSHTARIY', "G'ofurjonova Mushtariy"), ('Нўъмонжон кизи', "No'monjon qizi"),
    ("Jo'rayeva Zulfiya", "Jo'rayeva Zulfiya")])
def test_names(source, expected):
    assert latin_employee_name(source) == expected
    assert latin_employee_name(expected) == expected


def record(db, employee, key, amount):
    result = PayrollRecord(factory_code=employee.factory_code, employee_id=employee.id, dedupe_key=key, scan_uid=key,
        quantity=2, rate_per_piece=amount / 2, total_amount=amount, currency='UZS', scanned_at=datetime.now(timezone.utc))
    db.add(result)
    return result


def test_merge_preserves_ledger_managers_and_old_badges(client, auth_headers):
    with SessionLocal() as db:
        master = Employee(full_name='Aliyeva Mushtariy Voxidjon qizi', employee_no='322', factory_code='MIL', hr_profile_json={})
        target = Employee(full_name='Bir Xodim', employee_no='7011', factory_code='MIL')
        source = Employee(full_name='Bir Xodim', employee_no='7012', factory_code='MIL')
        db.add_all([master,target,source]); db.flush()
        worker = Employee(full_name='Another Worker', factory_code='MIL', manager_employee_id=source.id)
        db.add(worker); record(db,source,'merged-source',100); record(db,target,'merged-target',200); db.commit()
        source_id, target_id, master_id, worker_id = source.id, target.id, master.id, worker.id
        plan = {'names': [], 'merges': [{'source': source.id, 'target': target.id, 'factory':'MIL', 'name':'Bir Xodim', 'source_no':'7012', 'target_no':'7011'}], 'report_manager': master.id}
        result = reconcile_employees(db, plan); db.commit()
        assert result['merged_employees'][str(source_id)]['payroll_records.employee_id'] == 1
        assert db.get(Employee, source_id) is None
        assert db.get(Employee, worker_id).manager_employee_id == target_id
        rows=db.query(PayrollRecord).filter(PayrollRecord.employee_id==target_id).all()
        assert len(rows)==2 and sum(r.total_amount for r in rows)==300
        assert not db.query(SewingFlow).filter(SewingFlow.name.like('%Mushtariy%')).count()
        assert db.get(Employee,master_id).hr_profile_json['payroll_report_group']
    for url in [f'/api/payroll/qr/resolve/1{source_id:08}', '/api/payroll/employees/resolve?employee_no=7012', f'/api/payroll/employees/resolve?employee_no=EMP-{source_id}']:
        response=client.get(url,headers=auth_headers)
        assert response.status_code==200,response.text
        assert response.json()['employee_id']==target_id
    assert client.post('/api/employees',headers=auth_headers,json={'full_name':'Do not reuse','employee_no':'7012'}).status_code==409
    eco=_create_user_with_permissions(client,auth_headers,email='alias-eco@example.com',permissions=['payroll.scan'],factory_code='ECO')
    assert client.get(f'/api/payroll/qr/resolve/1{source_id:08}',headers=eco).status_code==404


def test_report_only_manager_group_and_export(client,auth_headers):
    with SessionLocal() as db:
        master=Employee(full_name='Aliyeva Mushtariy',factory_code='MIL',hr_profile_json={'payroll_report_group':True})
        db.add(master);db.flush()
        worker=Employee(full_name='Packaging Worker',factory_code='MIL',manager_employee_id=master.id)
        other=Employee(full_name='Other Worker',factory_code='MIL')
        db.add_all([worker,other]);db.flush()
        record(db,worker,'pack-work',240);record(db,other,'other-work',300)
        db.commit();mid=master.id;wid=worker.id
    options=client.get('/api/payroll/reports/sewing-production/options',headers=auth_headers)
    assert any(row['value']==f'manager:{mid}' for row in options.json()['sewing_lines'])
    url=f'/api/payroll/reports/sewing-production?sewing_line=manager:{mid}&report_view=salary'
    response=client.get(url,headers=auth_headers)
    assert response.status_code==200,response.text
    assert [row['employee_id'] for row in response.json()['salary_summary']]==[wid]
    exported=client.get(f'/api/payroll/reports/sewing-production.xlsx?sewing_line=manager:{mid}&report_view=salary',headers=auth_headers)
    assert exported.status_code==200,exported.text
    book=load_workbook(BytesIO(exported.content))
    assert 'Packaging Worker' in str(list(book.active.values))
    assert 'Other Worker' not in str(list(book.active.values))
    assert client.get('/api/payroll/reports/sewing-production?sewing_line=manager:bad',headers=auth_headers).status_code==400
    eco=_create_user_with_permissions(client,auth_headers,email='report-eco@example.com',permissions=['payroll.view'],factory_code='ECO')
    assert client.get(url,headers=eco).status_code==404


def test_branded_export_saved_sizes_literal_text_and_empty_order(client,auth_headers):
    with SessionLocal() as db:
        model=Model(code='TEST-V-6624',name='Template',details_json={'general':{'model_no':'PG10530','variant_no':'V-6624','qolip_no':'4388'}})
        order=BrandedPlanningOrder(order_no='BSO-EXPORT',ordered_for_type='milana',ordered_for_name='Milana',notes='=HYPERLINK("bad")')
        db.add_all([model,order]);db.flush()
        production=ProductionOrder(production_no='PO-EXPORT',production_type='branded_stock',planning_order_id=order.id,model_id=model.id,planned_quantity=600)
        db.add(production);db.flush()
        department=db.query(Department).filter_by(code='MIL').one()
        db.add(WorkOrder(production_order_id=production.id,department_id=department.id,operation='sewing',status='new'))
        db.add_all([ProductionOrderItem(production_order_id=production.id,model_id=model.id,color='Blue',size=size,planned_quantity=300,printing_required=False) for size in ['122','128']])
        db.commit();oid=order.id
    response=client.get(f'/api/planning/branded-orders/{oid}/export.xlsx',headers=auth_headers)
    assert response.status_code==200,response.text
    book=load_workbook(BytesIO(response.content))
    sheet=book['Buyurtma']
    assert sheet['B3'].value=='PG10530' and sheet['E3'].value=='V-6624' and sheet['D3'].value=='4388'
    assert sheet['L3'].value=='122, 128' and sheet['N3'].value==600
    assert sheet['O3'].value
    assert sheet.max_column==19
    assert sum(row[5] for row in list(book['Razmerlar'].values)[1:])==600
    assert all(cell.data_type!='f' for s in book for row in s for cell in row)
    assert client.get('/api/planning/branded-orders/999999/export.xlsx',headers=auth_headers).status_code==404
    restricted=_create_user_with_permissions(client,auth_headers,email='no-export@example.com',permissions=['payroll.view'])
    assert client.get(f'/api/planning/branded-orders/{oid}/export.xlsx',headers=restricted).status_code==403


def test_excel_embeds_webp_as_png_without_remote_fetch():
    import base64
    from zipfile import ZipFile
    from PIL import Image
    from openpyxl import Workbook
    from app.models.branded_order_export_service import picture
    image=BytesIO()
    Image.new('RGB',(20,30),(50,100,150)).save(image,format='WEBP')
    book=Workbook()
    picture(book.active,'C3','data:image/webp;base64,'+base64.b64encode(image.getvalue()).decode())
    picture(book.active,'C4','https://example.invalid/no-network-request.webp')
    output=BytesIO();book.save(output)
    with ZipFile(output) as archive:
        media=[name for name in archive.namelist() if name.startswith('xl/media/')]
        assert len(media)==1 and media[0].endswith('.png')
    loaded=load_workbook(BytesIO(output.getvalue()))
    assert loaded.active['C4'].value=='https://example.invalid/no-network-request.webp'
