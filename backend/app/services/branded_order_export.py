"""Read-only branded-order workbook matching the owner's 19-column production form."""
from io import BytesIO
import base64
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PillowImage
from sqlalchemy.orm import selectinload

from app.models import BrandedPlanningOrder, Department, Model, ProductionOrder, ProductionOrderMaterial, StockBatch
from app.services.model_identity import model_number_fields
from app.services.label_images import model_label_image_src, material_label_image_src

HEADERS = ['No', 'Model №', 'Rasm', 'Model qolip №', 'Variant №', 'Material', 'Material eni', 'Rangi', 'Partiya №', 'Pechat', 'Vishevka', 'Model razmeri', 'Gramaj material', 'Umumiy ish soni', 'Potok', 'Mini postal', 'Aksesuar', 'Izoh', 'Foto izoh']


def append(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if isinstance(cell.value, (int, float, Decimal)):
            cell.number_format = '#,##0' if float(cell.value).is_integer() else '#,##0.######'


def first(values, *keys):
    return next((values[key] for key in keys if values.get(key) not in (None, '')), None)


def picture(sheet, cell, source):
    # Embedded uploads only: never fetch an arbitrary URL from a model field.
    if not source or not source.startswith('data:image/'):
        if source:
            sheet[cell] = source
            sheet[cell].data_type = 's'
        return
    raw = base64.b64decode(source.split(',', 1)[1])
    try:
        # Excel/OpenXML does not support WebP. Normalize uploaded formats to PNG.
        with PillowImage.open(BytesIO(raw)) as original:
            original.thumbnail((540, 480))
            normalized = BytesIO()
            original.convert('RGBA').save(normalized, format='PNG')
        normalized.seek(0)
        image = Image(normalized)
        ratio = min(135 / image.width, 120 / image.height, 1)
        image.width *= ratio
        image.height *= ratio
        sheet.add_image(image, cell)
    except (OSError, ValueError):
        sheet[cell] = 'Image unavailable'


def build_branded_order_workbook(db, order: BrandedPlanningOrder):
    productions = db.query(ProductionOrder).options(
        selectinload(ProductionOrder.items), selectinload(ProductionOrder.work_orders),
        selectinload(ProductionOrder.materials).selectinload(ProductionOrderMaterial.stock_batch),
    ).filter(ProductionOrder.planning_order_id == order.id).order_by(ProductionOrder.id).all()
    models = {m.id: m for m in db.query(Model).options(selectinload(Model.images), selectinload(Model.bom)).filter(
        Model.id.in_([p.model_id for p in productions])).all()}
    batches = {b.id: b for b in db.query(StockBatch).filter(StockBatch.id.in_([
        p.fabric_batch_id for p in productions if p.fabric_batch_id])).all()}
    departments = dict(db.query(Department.id, Department.name).filter(Department.id.in_(
        {w.department_id for p in productions for w in p.work_orders})).all())
    book = Workbook()
    form = book.active
    form.title = 'Buyurtma'
    form.merge_cells('A1:S1')
    append(form, HEADERS)
    form['A1'] = f'Milana Textile — {order.order_no} — {order.ordered_for_name} — {order.created_at:%d.%m.%Y}'
    form['A1'].data_type = 's'
    form['A1'].font = Font(size=16, bold=True)
    form.row_dimensions[1].height = 30
    sizes = book.create_sheet('Razmerlar')
    append(sizes, ['Buyurtma', 'Model', 'Variant', 'Rang', 'Razmer', 'Soni', 'Pechat'])
    materials = book.create_sheet('Materiallar')
    append(materials, ['Buyurtma', 'Model', 'Variant', 'Material', 'Partiya', 'Eni', 'Gramaj', 'Rang', 'Reja sarfi', 'Birlik'])
    details = book.create_sheet('Tafsilotlar')
    append(details, ['Buyurtma', 'Maydon', 'Qiymat'])
    for key, value in [('Status', order.status), ('Izoh', order.notes)]:
        append(details, [order.order_no, key, value])
    for index, p in enumerate(productions, 1):
        model = models.get(p.model_id)
        identity = model_number_fields(model)
        general = (model.details_json or {}).get('general', {}) if model else {}
        general = general if isinstance(general, dict) else {}
        fabric_rows = [(m.stock_batch, m.estimated_quantity, m.unit) for m in p.materials]
        if not fabric_rows and p.fabric_batch_id in batches:
            fabric_rows = [(batches[p.fabric_batch_id], p.estimated_material_amount, p.estimated_material_unit)]
        def joined(field):
            return '\n'.join(str(getattr(batch, field) or '') for batch, _, _ in fabric_rows) or None
        fabric = '\n'.join(batch.item.name for batch, _, _ in fabric_rows)
        if not fabric and model:
            fabric = '\n'.join(b.material_name or (b.item.name if b.item else '') for b in model.bom
                               if b.item and b.item.category in {'fabric', 'semi_finished'})
        accessories = '\n'.join(f'{b.material_name or (b.item.name if b.item else "")} — {b.quantity_per_piece} {b.unit}'
                                for b in model.bom if b.item and b.item.category in {'accessory', 'packaging'}) if model else ''
        sewing = [departments[w.department_id] for w in p.work_orders if w.operation == 'sewing' and w.department_id in departments]
        notes = '\n'.join(str(v) for v in [p.order_no, p.printing_instructions, order.notes] if v)
        append(form, [index, identity['model_no'], None,
                      first(general, 'qolip_no', 'qolipNo', 'mold_no', 'moldNo', 'pattern_no', 'patternNo'),
                      identity['variant_no'], fabric, joined('width'), joined('color'), joined('batch_no'),
                      'Ha' if any(i.printing_required for i in p.items) else 'Yo‘q',
                      ('Ha' if first(general, 'embroidery', 'legacy_master_embroidery') is True else
                       'Yo‘q' if first(general, 'embroidery', 'legacy_master_embroidery') is False else
                       first(general, 'embroidery', 'legacy_master_embroidery')),
                      ', '.join(dict.fromkeys(i.size for i in p.items if i.size)), joined('gsm'), p.planned_quantity,
                      ', '.join(dict.fromkeys(sewing)), first(general, 'mini_postal'), accessories, notes, None])
        row = form.max_row
        form.row_dimensions[row].height = 120
        picture(form, f'C{row}', model_label_image_src(model))
        picture(form, f'H{row}', material_label_image_src(model))
        form[f'H{row}'].alignment = Alignment(vertical='bottom', wrap_text=True)
        for item in p.items:
            append(sizes, [p.order_no, identity['model_no'], identity['variant_no'], item.color, item.size, item.planned_quantity, item.printing_required])
        for batch, amount, unit in fabric_rows:
            append(materials, [p.order_no, identity['model_no'], identity['variant_no'], batch.item.name, batch.batch_no,
                               batch.width, batch.gsm, batch.color, amount, unit])
        for key, value in [('Status', p.status), ('Boshlanish', p.start_date), ('Muddat', p.deadline),
                           ('Pechat izohi', p.printing_instructions), ('Reja soni', p.planned_quantity)]:
            append(details, [p.order_no, key, value.isoformat() if hasattr(value, 'isoformat') else value])
        for attachment in p.printing_attachments or []:
            for key, value in attachment.items():
                append(details, [p.order_no, f'Pechat fayli: {key}', str(value)])
        for key, value in general.items():
            append(details, [p.order_no, f'Model: {key}', str(value) if value is not None else None])
    border = Border(*(Side(style='thin', color='CCCCCC') for _ in range(4)))
    for sheet in book:
        header_row = 2 if sheet == form else 1
        sheet.freeze_panes = f'C{header_row + 1}'
        sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        for row in sheet.iter_rows(min_row=header_row):
            for cell in row:
                cell.border = border
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='E7E6E1')
        sheet.row_dimensions[header_row].height = 36
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 22
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = f'1:{header_row}'
    form.column_dimensions['A'].width = 6
    form.column_dimensions['R'].width = 35
    details.column_dimensions['C'].width = 65
    output = BytesIO()
    book.save(output)
    return output.getvalue()
