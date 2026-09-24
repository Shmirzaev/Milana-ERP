"""Read-only Excel version of the same shipment invoice used for printing."""
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from textwrap import wrap

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.shipment_invoice import LABELS, WEIGHT_NOTE, build_invoice_rows, recorded_weight


def shipment_invoice_workbook(document: dict, language: str) -> bytes:
    lang = language if language in LABELS else "en"
    text = LABELS[lang]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoice"
    widths = [6, 14, 13, 30, 34, 10, 12, 13, 13]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    def put(row, column, value, *, bold=False, fill=None, color="243446", align="left"):
        cell = sheet.cell(row, column, value)
        # Business text remains literal even when it starts with =, +, - or @.
        if isinstance(value, str):
            cell.data_type = "s"
        cell.font = Font(name="Calibri", size=10, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        return cell

    def merged(row, first, last, value, **style):
        sheet.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
        put(row, first, value, **style)

    merged(1, 1, 9, "MILANA PREMIUM", bold=True, color="9B242B")
    merged(2, 1, 9, f'{text["title"]} · {document["shipment_no"]}', bold=True, fill="243446", color="FFFFFF")
    sheet.row_dimensions[2].height = 28
    transport = document.get("transport_details") or {}
    shipped = document.get("shipped_at")
    if shipped:
        try:
            parsed = datetime.fromisoformat(str(shipped).replace("Z", "+00:00"))
            shipped = (parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)).astimezone(timezone(timedelta(hours=5))).strftime("%d/%m/%Y %H:%M:%S")
        except ValueError:
            pass
    metadata = [
        (text["customer"], document.get("customer"), text["driver"], transport.get("driver_name")),
        (text["order"], document.get("sales_order_no"), text["vehicle"], transport.get("vehicle_info")),
        (text["date"], shipped, text["carrier"], transport.get("cargo_name")),
        (text["issued"], document.get("warehouse_person"), text["phone"], transport.get("driver_phone")),
    ]
    for row, (left, lv, right, rv) in enumerate(metadata, 4):
        merged(row, 1, 5, f'{left}: {lv or "—"}')
        merged(row, 6, 9, f'{right}: {rv or "—"}')
        sheet.row_dimensions[row].height = 30
    headers = ["№", text["modelNo"], text["variant"], text["description"], text["size"], text["packs"], text["qty"], text["weight"], text["totalWeight"]]
    for column, title in enumerate(headers, 1):
        put(9, column, title, bold=True, fill="243446", color="FFFFFF", align="center")
    sheet.row_dimensions[9].height = 30
    packages = document.get("package_details")
    if packages is None:
        packages = list({line["package_no"]: {"package_no": line["package_no"], "weight_kg": None} for line in document.get("lines", [])}.values())
    rows = document.get("invoice_rows") or build_invoice_rows(document.get("lines", []), packages)
    border = Border(bottom=Side(style="hair", color="D5DCE2"))
    for index, item in enumerate(rows, 1):
        row = index + 9
        sizes = item.get("sizes", [])
        multicolor = len({s.get("color") for s in sizes if s.get("color")}) > 1
        size_text = "\n".join((f'{s.get("color") or ""} / ' if multicolor else "") + str(s.get("size") or "") +
                              ("" if s.get("aggregate") else f' ({int(s["quantity"])})') for s in sizes)
        span = item["package_rowspan"]
        weight = float(Decimal(str(item["weight_kg"]))) if item.get("weight_kg") is not None else None
        values = [index, item.get("model_no"), item.get("variant_no"), item.get("description"), size_text,
                  item["pack_count"] if span else None, item["quantity"], weight, weight]
        for column, value in enumerate(values, 1):
            cell = put(row, column, value, fill="F3F6F8" if index % 2 == 0 else "FFFFFF", align="right" if column >= 6 else "left")
            cell.border = border
            if column >= 6:
                cell.number_format = '#,##0.00' if column >= 8 else '#,##0'
        for column in (6, 8, 9):
            if span > 1:
                sheet.merge_cells(start_row=row, end_row=row + span - 1, start_column=column, end_column=column)
        line_count = max(sum(max(1, len(wrap(part, max(1, int(width - 2))))) for part in str(value or "").split("\n")) for value, width in zip(values, widths))
        sheet.row_dimensions[row].height = max(23, line_count * 13 + 8)
    total_row = 10 + len(rows)
    merged(total_row, 1, 5, text["total"], bold=True, fill="E8F1EC", color="174A35")
    for column, value in enumerate([document["packages_count"], document["quantity"], float(recorded_weight(document)), float(recorded_weight(document))], 6):
        cell = put(total_row, column, value, bold=True, fill="E8F1EC", color="174A35", align="right")
        cell.number_format = '#,##0.00' if column >= 8 else '#,##0'
    sheet.row_dimensions[total_row].height = 27
    notes = []
    if document.get("missing_weight_packages", sum(p.get("weight_kg") is None for p in packages)):
        notes.append(WEIGHT_NOTE[lang])
    if document.get("historical_reconstruction") or document.get("invoice_layout_version") != 2:
        notes.append(text["historical"])
    notes.append(text["ledger"] + ": " + str(document.get("ledger_invoice_no") or "") if document.get("finance_posting_status") == "posted" else text.get(document.get("finance_posting_status"), text["posting"]))
    for row, note in enumerate(notes, total_row + 2):
        merged(row, 1, 9, note, color="575E66")
        sheet.row_dimensions[row].height = 30
    sheet.freeze_panes = "D10"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "9:9"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:I{sheet.max_row}"
    sheet.oddFooter.right.text = "&P / &N"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
