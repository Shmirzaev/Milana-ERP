"""Read-only full ready-product warehouse workbook."""

from datetime import datetime, timedelta, timezone
from io import BytesIO
from textwrap import wrap

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import func
from sqlalchemy.orm import load_only, raiseload, selectinload

from app.models import Model, ModelColor, Package
from app.services.model_identity import model_number_fields


TEXT = {
    "en": ["Ready product warehouse", "Standard stock", "First Grade", "Model number", "Variant number", "Packages", "Pieces", "Total", "All warehouse stock, including reserved and unplaced packages. Page filters do not apply. First Grade packages are individual pieces."],
    "ru": ["Склад готовой продукции", "Основной склад", "Первый сорт", "Номер модели", "Номер варианта", "Упаковки", "Штуки", "Итого", "Все складские остатки, включая резерв и упаковки без места. Фильтры страницы не применяются. Упаковки первого сорта — отдельные изделия."],
    "uz": ["Tayyor mahsulotlar ombori", "Asosiy ombor", "Birinchi sort", "Model raqami", "Variant raqami", "Qop soni", "Dona soni", "Jami", "Barcha ombor qoldiqlari, jumladan band qilingan va joylashtirilmagan qoplar. Sahifa filtrlari qo‘llanmaydi. Birinchi sort qadoqlari alohida donalardir."],
}

DETAIL_TEXT = {
    "en": ["Product name", "Color", "Exported at (Tashkent time, UTC+05:00)", "Inventory source"],
    "ru": ["Название товара", "Цвет", "Дата и время экспорта (время Ташкента, UTC+05:00)", "Источник остатков"],
    "uz": ["Mahsulot nomi", "Rang", "Eksport vaqti (Toshkent vaqti, UTC+05:00)", "Qoldiq manbasi"],
}


def model_colors(model: Model | None) -> str:
    # Keep the existing model/variant grain, including when several colors are
    # recorded. Never join this one-to-many relationship into the stock totals.
    names = dict.fromkeys(
        color.color_name for color in sorted(model.colors, key=lambda color: color.id)
        if color.color_name and color.color_name.strip()
    ) if model else {}
    return "\n".join(names) or "—"


def row_height(values, widths) -> float:
    # Excel does not auto-fit wrapped rows reliably on opening generated files.
    lines = max(
        sum(max(1, len(wrap(line, width=max(1, int(width * 0.85))))) for line in str(value or "").split("\n"))
        for value, width in zip(values, widths)
    )
    return max(26, lines * 16 + 10)


def export_warehouse_report(db, lang: str) -> bytes:
    # Match Warehouse Stock's ready statuses, including imported and unplaced stock.
    quantities = (
        db.query(Package.stock_kind, Package.model_id, func.count(Package.id), func.sum(Package.total_quantity))
        .filter(Package.status.in_(["packed", "received_in_storage", "reserved"]))
        .group_by(Package.stock_kind, Package.model_id)
        .all()
    )
    model_ids = {row.model_id for row in quantities if row.model_id is not None}
    models = {
        model.id: model
        for model in db.query(Model).options(
            raiseload("*"), load_only(Model.id, Model.code, Model.name, Model.details_json),
            selectinload(Model.colors).load_only(ModelColor.id, ModelColor.model_id, ModelColor.color_name),
        ).filter(Model.id.in_(model_ids)).all()
    } if model_ids else {}
    labels = TEXT[lang]
    detail = DETAIL_TEXT[lang]
    workbook = Workbook()
    workbook.remove(workbook.active)
    generated = datetime.now(timezone(timedelta(hours=5))).strftime("%d.%m.%Y %H:%M:%S")
    widths = {"A": 24, "B": 52, "C": 24, "D": 30, "E": 16, "F": 18}
    thin = Side(style="thin", color="D9DDE3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for kind, title in [("standard", labels[1]), ("first_grade", labels[2])]:
        sheet = workbook.create_sheet(title)
        metadata = [labels[0], f"{detail[2]}: {generated}",
                    f"{detail[3]}: ERP — {labels[0]} — {title}", labels[8]]
        for number, value in enumerate(metadata, 1):
            sheet.merge_cells(start_row=number, start_column=1, end_row=number, end_column=6)
            cell = sheet.cell(number, 1, value)
            cell.font = Font(name="Calibri", size=16 if number == 1 else 11, bold=number == 1, color="20252B")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            sheet.row_dimensions[number].height = 32 if number == 1 else 42 if number == 4 else 24
        sheet.append([labels[3], detail[0], labels[4], detail[1], labels[5], labels[6]])
        rows = [row for row in quantities if row.stock_kind == kind]
        rows.sort(key=lambda row: tuple(str(value or "").casefold() for value in model_number_fields(models.get(row.model_id)).values()))
        for index, (_, model_id, packages, pieces) in enumerate(rows):
            model = models.get(model_id)
            identity = model_number_fields(model)
            values = [identity["model_no"], model.name if model else None, identity["variant_no"],
                      model_colors(model), int(packages), int(pieces or 0)]
            sheet.append(values)
            row_number = sheet.max_row
            sheet.row_dimensions[row_number].height = row_height(values, widths.values())
            for cell in sheet[row_number]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                cell.font = Font(name="Calibri", size=11, color="20252B")
                cell.fill = PatternFill("solid", fgColor="F5F6F8" if index % 2 else "FFFFFF")
                cell.border = border
                cell.alignment = Alignment(horizontal="right" if cell.column >= 5 else "left", vertical="center", wrap_text=True)
                if isinstance(cell.value, int):
                    cell.number_format = "#,##0"
        last = sheet.max_row
        sheet.append([labels[7], None, None, None,
                      f"=SUM(E6:E{last})" if rows else 0,
                      f"=SUM(F6:F{last})" if rows else 0])
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=4)
        for number in (5, sheet.max_row):
            for cell in sheet[number]:
                cell.font = Font(name="Calibri", size=11, bold=True, color="20252B")
                cell.fill = PatternFill("solid", fgColor="E9EDF1" if number == 5 else "E1E6EB")
                cell.border = border if number == 5 else Border(left=thin, right=thin, bottom=thin, top=Side(style="medium", color="9CA3AF"))
                cell.alignment = Alignment(horizontal="right" if cell.column >= 5 else "left", vertical="center", wrap_text=True)
                if cell.column >= 5:
                    cell.number_format = "#,##0"
            sheet.row_dimensions[number].height = 30
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A5:F{last}"
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = "1:5"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
