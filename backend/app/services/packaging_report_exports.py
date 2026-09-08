"""Excel export using the application's existing server-side workbook stack."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image


TEXT = {
    "en": {
        "entries": "Packaging output",
        "completed": "Completed jobs",
        "daily": "Daily totals",
        "date": "Date",
        "passport": "Order passports",
        "order_no": "Production order",
        "batch": "Batch",
        "brand": "Brand",
        "model_no": "Model",
        "variant_no": "Variant",
        "category": "Category",
        "color": "Color",
        "sizes": "Sizes",
        "standards": "Standard qty",
        "package_count": "Packages",
        "quantity": "Quantity",
        "two_piece_quantity": "2x2 quantity",
        "first_sort": "1st grade",
        "second_sort": "2nd grade",
        "cutting_defects": "Cutting defects",
        "planned_quantity": "Planned qty",
        "packed_quantity": "Packed output",
        "damaged_quantity": "Damaged",
        "shortage": "Shortage",
        "balance": "Balance",
        "notes": "Notes",
        "total": "Total",
        "image": "Picture",
        "note": "Package creation dates use Tashkent time. Blank grade, 2x2, cutting defect and shortage fields are not recorded separately. Completed jobs use completion dates and lifetime workflow totals; balance = packed + damaged - planned. Passports refer to the whole order.",
    },
    "ru": {
        "entries": "Выпуск упаковки",
        "completed": "Завершённые работы",
        "daily": "Итоги по дням",
        "date": "Дата",
        "passport": "Паспорта заказа",
        "order_no": "Производственный заказ",
        "batch": "Партия",
        "brand": "Бренд",
        "model_no": "Модель",
        "variant_no": "Вариант",
        "category": "Категория",
        "color": "Цвет",
        "sizes": "Размеры",
        "standards": "Стандарт шт.",
        "package_count": "Кол-во упаковок",
        "quantity": "Количество",
        "two_piece_quantity": "Шт. 2x2",
        "first_sort": "1-сорт",
        "second_sort": "2-сорт",
        "cutting_defects": "Брак кроя",
        "planned_quantity": "План шт.",
        "packed_quantity": "Упаковано",
        "damaged_quantity": "Повреждено",
        "shortage": "Недостача",
        "balance": "Баланс",
        "notes": "Примечания",
        "total": "Итого",
        "image": "Фото",
        "note": "Даты создания упаковок указаны по Ташкенту. Пустые поля сорта, 2x2, брака кроя и недостачи отдельно не учитываются. Завершённые работы отбираются по дате завершения; объёмы — за всё время работы. Баланс = упаковано + повреждено - план. Паспорта относятся ко всему заказу.",
    },
    "uz": {
        "entries": "Upakovka ishlari",
        "completed": "Yopilgan ishlar",
        "daily": "Kunlik jami",
        "date": "Sana",
        "passport": "Buyurtma pasportlari",
        "order_no": "Ishlab chiqarish buyurtmasi",
        "batch": "Partiya",
        "brand": "Brend",
        "model_no": "Model",
        "variant_no": "Variant",
        "category": "Kategoriya",
        "color": "Rang",
        "sizes": "Razmer",
        "standards": "Standart dona",
        "package_count": "Qop soni",
        "quantity": "Ish soni",
        "two_piece_quantity": "2x2 ish soni",
        "first_sort": "1-sort",
        "second_sort": "2-sort",
        "cutting_defects": "Kroy brak",
        "planned_quantity": "Reja soni",
        "packed_quantity": "Upakovka ish soni",
        "damaged_quantity": "Shikastlangan",
        "shortage": "Kamomad",
        "balance": "Balans",
        "notes": "Izoh",
        "total": "Umumiy",
        "image": "Rasm",
        "note": "Qop yaratilgan sana Toshkent vaqti bilan olinadi. Bo‘sh sort, 2x2, kroy brak va kamomad maydonlari alohida qayd etilmagan. Yopilgan ishlar yakunlangan sana bo‘yicha olinadi; miqdorlar ishning jami natijasidir. Balans = upakovka + shikastlangan - reja. Pasportlar butun buyurtmaga tegishli.",
    },
}

ENTRY_COLUMNS = [
    "date",
    "passport",
    "order_no",
    "batch",
    "brand",
    "model_no",
    "variant_no",
    "category",
    "color",
    "sizes",
    "standards",
    "package_count",
    "quantity",
    "two_piece_quantity",
    "image",
    "first_sort",
    "second_sort",
]
COMPLETED_COLUMNS = [
    "date",
    "passport",
    "order_no",
    "batch",
    "brand",
    "model_no",
    "variant_no",
    "category",
    "planned_quantity",
    "packed_quantity",
    "damaged_quantity",
    "first_sort",
    "second_sort",
    "shortage",
    "balance",
    "notes",
]
DAILY_COLUMNS = [
    "date",
    "package_count",
    "quantity",
    "two_piece_quantity",
    "first_sort",
    "second_sort",
    "cutting_defects",
]


def export_packaging_report(report: dict, lang: str) -> bytes:
    labels = TEXT[lang]
    workbook = Workbook()
    workbook.remove(workbook.active)
    factory = {"PKG": "Milana", "BPK": "Besttex", "ECP": "Eco Cotton"}[report["packaging_department_code"]]
    numeric = {
        "package_count",
        "quantity",
        "two_piece_quantity",
        "first_sort",
        "second_sort",
        "cutting_defects",
        "planned_quantity",
        "packed_quantity",
        "damaged_quantity",
        "shortage",
        "balance",
    }
    for kind, columns in [("entries", ENTRY_COLUMNS), ("completed", COMPLETED_COLUMNS), ("daily", DAILY_COLUMNS)]:
        sheet = workbook.create_sheet(labels[kind])
        count = len(columns) + 1
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=count)
        sheet.cell(1, 1, f"{factory} — {labels[kind]} | {report['from_date']} – {report['to_date']}")
        sheet.cell(1, 1).font = Font(size=15, bold=True)
        sheet.row_dimensions[1].height = 28
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=count)
        sheet.cell(2, 1, labels["note"]).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[2].height = 65
        sheet.append(["№", *[labels[c] for c in columns]])
        for cell in sheet[3]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E9E5DA")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[3].height = 34
        for index, row in enumerate(report[kind], 1):
            sheet.append([index, *[date.fromisoformat(row[c]) if c == "date" else row.get(c) for c in columns]])
            for cell in sheet[sheet.max_row]:
                # Business text must never be interpreted as an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color="D8D3C7"))
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, int):
                    cell.number_format = "#,##0"
            sheet.row_dimensions[sheet.max_row].height = 32
            if kind == "entries" and row.get("_image"):
                picture = Image(BytesIO(row["_image"]))
                sheet.add_image(picture, f"{get_column_letter(columns.index('image') + 2)}{sheet.max_row}")
                sheet.row_dimensions[sheet.max_row].height = 72
        end_row = sheet.max_row
        total_row = end_row + 1
        sheet.cell(total_row, 1, labels["total"])
        for column_index, key in enumerate(columns, 2):
            if key in numeric and any(row.get(key) is not None for row in report[kind]):
                letter = get_column_letter(column_index)
                sheet.cell(total_row, column_index, f"=SUM({letter}4:{letter}{end_row})").number_format = "#,##0"
        for cell in sheet[total_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F1EEE6")
        sheet.column_dimensions["A"].width = 10
        for column_index, key in enumerate(columns, 2):
            sheet.column_dimensions[get_column_letter(column_index)].width = (
                32 if key in {"category", "notes", "passport", "order_no", "batch"} else 18
            )
        sheet.freeze_panes = "F4" if kind != "daily" else "C4"
        sheet.auto_filter.ref = f"A3:{get_column_letter(count)}{max(3, end_row)}"
        sheet.print_title_rows = "1:3"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if kind != "daily" else sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
        sheet.print_area = f"A1:{get_column_letter(count)}{total_row}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
