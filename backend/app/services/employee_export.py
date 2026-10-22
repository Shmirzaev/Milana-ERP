"""Factory employee directory export, with the same field visibility as the list."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


LABELS = {
    "en": ("Employees", ["Employee ID", "Full name", "Factory", "Department", "Position", "Manager", "Status", "Hire date", "Phone"],
           {"active": "Active", "inactive": "Inactive", "on_leave": "On leave", "terminated": "Terminated", "archived": "Archived"}),
    "ru": ("Сотрудники", ["Номер сотрудника", "ФИО", "Фабрика", "Отдел", "Должность", "Руководитель", "Статус", "Дата приёма", "Телефон"],
           {"active": "Активен", "inactive": "Неактивен", "on_leave": "В отпуске", "terminated": "Уволен", "archived": "В архиве"}),
    "uz": ("Xodimlar", ["Xodim raqami", "F.I.Sh.", "Fabrika", "Bo‘lim", "Lavozim", "Rahbar", "Holat", "Ishga kirgan sana", "Telefon"],
           {"active": "Faol", "inactive": "Nofaol", "on_leave": "Ta’tilda", "terminated": "Ishdan bo‘shatilgan", "archived": "Arxivda"}),
}


def employee_workbook(rows, lang="en", *, include_private=False):
    title, headers, statuses = LABELS[lang]
    book = Workbook(write_only=True)
    sheet = book.create_sheet(title)
    sheet.freeze_panes = "C2"
    headers = headers if include_private else headers[:-1]
    for index, width in enumerate([20, 36, 14, 26, 26, 36, 20, 20, 22][:len(headers)], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    def append(values, header=False):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # Preserve leading zeros and treat user content as literal text.
            if header:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="475569")
            cells.append(cell)
        sheet.append(cells)

    append(headers, header=True)
    count = 1
    for employee, department, position, manager in rows:
        values = [str(employee.employee_no or "").strip() or str(employee.id).zfill(4), employee.full_name,
                  employee.factory_code, department, position or employee.position, manager,
                  statuses.get(employee.status, employee.status), employee.joined_at.date().isoformat() if employee.joined_at else None]
        if include_private:
            values.append(employee.phone)
        append(values)
        count += 1
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
