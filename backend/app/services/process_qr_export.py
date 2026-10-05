"""Read-only Excel rendering of the Process QR editor's current operation list."""
from decimal import Decimal
from io import BytesIO
from typing import Annotated, Literal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel, Field


Text = Annotated[str, Field(max_length=500)]


class ProcessQrExportRow(BaseModel):
    selected: bool
    section: Text
    code: Text
    name: Text
    rate: Decimal = Field(ge=0, le=Decimal("9999999999.9999"), allow_inf_nan=False)
    copies: int = Field(ge=1, le=10000)
    division: Text


class ProcessQrExportIn(BaseModel):
    model: Text
    order: Text = ""
    factory: Literal["MIL", "BST", "ECO"]
    currency: Annotated[str, Field(max_length=10)]
    headers: list[Text] = Field(min_length=8, max_length=8)
    rows: list[ProcessQrExportRow] = Field(min_length=1, max_length=1000)


def build_process_qr_xlsx(payload: ProcessQrExportIn) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Paid processes"
    def append(values):
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value for value in values])

    append([payload.model, payload.order, payload.factory, payload.currency])
    append(payload.headers)
    for index, row in enumerate(payload.rows, 1):
        append([index, row.selected, row.section, row.code, row.name, row.rate, row.copies, row.division])
        sheet.cell(index + 2, 6).number_format = "#,##0.00##"
    # All user-entered strings are literal cells, never spreadsheet formulas.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[2]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGH", [24, 24, 22, 22, 52, 20, 12, 32]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:H{sheet.max_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
