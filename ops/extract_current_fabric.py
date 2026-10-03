"""Read source workbooks, including binary SAMO, without rewriting originals."""
import hashlib
import json
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/fabric-reconciliation'
sys.path.insert(0, str(OUT / 'python-deps'))

FILES = ['DINAR 2026 (5) (6).xlsx', 'SAMO 2025 (4) (2) (version 1) (4).xlsb',
         'Cафф милана  2025 (4) (2) (2) (4).xlsx']


def extract(files=None):
    result = []
    for name in files or FILES:
        path = Path('C:/Users/User/Downloads/Telegram Desktop') / name
        book = dict(file=name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(), sheets=[])
        if path.suffix == '.xlsb':
            from pyxlsb import open_workbook
            with open_workbook(str(path)) as wb:
                for title in wb.sheets:
                    rows = []
                    with wb.get_sheet(title) as sheet:
                        for cells in sheet.rows(sparse=True):
                            values = [None] * max(21, max((c.c for c in cells), default=0)+1)
                            for c in cells:
                                values[c.c] = int(c.v) if isinstance(c.v, float) and c.v.is_integer() else c.v
                            if any(v is not None for v in values):
                                rows.append(dict(row=cells[0].r+1, values=values, formulas={}))
                    book['sheets'].append(dict(sheet=title, rows=rows))
        else:
            with path.open('rb') as stream:
                wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
                fm = openpyxl.load_workbook(path, read_only=True, data_only=False)
                for sheet in wb:
                    rows = []
                    with zipfile.ZipFile(path) as archive:
                        ns = {'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                        xml = ET.fromstring(archive.read(sheet._worksheet_path))
                        populated = [c.attrib['r'] for c in xml.findall('.//m:sheetData/m:row/m:c',ns)
                                     if c.find('m:v',ns) is not None or c.find('m:f',ns) is not None or c.find('m:is',ns) is not None]
                    bounds = [openpyxl.utils.cell.coordinate_to_tuple(c) for c in populated]
                    last_row = max((r for r,c in bounds),default=1)
                    last_col = max(21,max((c for r,c in bounds),default=1))
                    for n, (values, formulas) in enumerate(zip(sheet.iter_rows(max_row=last_row,max_col=last_col,values_only=True), fm[sheet.title].iter_rows(max_row=last_row,max_col=last_col,values_only=True)), 1):
                        if not any(v is not None for v in formulas):
                            continue
                        rows.append(dict(row=n, values=list(values)+[None]*max(0,21-len(values)),
                            formulas={openpyxl.utils.get_column_letter(i+1):v for i,v in enumerate(formulas) if isinstance(v,str) and v.startswith('=')}))
                    book['sheets'].append(dict(sheet=sheet.title, rows=rows))
                wb.close(); fm.close()
        result.append(book)
    (OUT/'source-workbooks.json').write_text(json.dumps(result, ensure_ascii=False, default=str), encoding='utf-8')
    print(json.dumps([dict(file=b['file'], sheets=[(s['sheet'],len(s['rows'])) for s in b['sheets']]) for b in result]))


if __name__ == '__main__':
    extract(sys.argv[1:] or None)
