"""Extract bounded workbook evidence, treating all cell content as data."""
import hashlib
import json
from pathlib import Path
import openpyxl

SOURCE = Path('C:/Users/User/Downloads/Telegram Desktop')
FILES = ['DINAR 2026 (5).xlsx', 'SAMO 2025 (4) (2).xlsx', 'Cафф милана  2025 (4) (2) (2).xlsx']
OUT = Path(__file__).resolve().parents[1] / 'outputs/fabric-reconciliation'

def extract():
    evidence = []
    for filename in FILES:
        path = SOURCE / filename
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        book = {'file':filename, 'sha256':hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets':[]}
        for sheet in wb:
            # These supplied sheets have a verified last populated row <= 579.
            rows = []
            for n, (values, source) in enumerate(zip(sheet.iter_rows(max_row=1100,max_col=21,values_only=True), formulas[sheet.title].iter_rows(max_row=1100,max_col=21,values_only=True)),1):
                if not any(v is not None for v in source):
                    continue
                rows.append({'row':n, 'values':list(values), 'formulas':{openpyxl.utils.get_column_letter(i+1):v for i,v in enumerate(source) if isinstance(v,str) and v.startswith('=')}})
            book['sheets'].append({'sheet':sheet.title,'rows':rows})
        evidence.append(book)
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/'source-workbooks.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    return evidence

if __name__ == '__main__':
    evidence=extract()
    print(json.dumps([{ 'file':b['file'],'sheets':[(s['sheet'],len(s['rows'])) for s in b['sheets']]} for b in evidence],ensure_ascii=True))
